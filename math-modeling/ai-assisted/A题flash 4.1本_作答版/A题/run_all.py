# -*- coding: utf-8 -*-
"""
生成 result1~result4.xlsx 及论文表1~表6  （v2 优化版）

关键设定
--------
* 网格 N=1281（经网格收敛检验；v1 的 N=141 对应表面值误差约 3e-3，未收敛；
  N=321/641/1281 的收敛序列见 解答报告.html 第 3.4 节）
* 时间积分 BDF 自适应步长，rtol=1e-7；已与"后向欧拉 Δt→0"极限和
  MATLAB 独立实现(pdepe)交叉验证
* 问题4 在 Landau 坐标 xi=r/R(t) 下求解，【不含对流项】（仿射收缩下相对速度为零），
  输出采用【物质坐标】r0 = xi*R0（结果模板末列为"药材表面"）
* 结果保留 4 位小数
"""
import numpy as np
import openpyxl
import os
import solver as sv

BASE = sv.BASE
OUTDIR = os.path.join(BASE, "结果")
os.makedirs(OUTDIR, exist_ok=True)

N = 1281
DC = np.round(np.arange(0, 2.0 + 1e-9, 0.1), 1)      # 0.0 .. 2.0 共 21 列(物质坐标)
TAB_D = [0, 0.5, 1.0, 1.5, 2.0]


def write_grid(path, sheets, header, times):
    """sheets: {表名: 2D 数组 (nt,nd)}; header: 表头(nd); times: 时间列(nt)"""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sn, arr in sheets.items():
        ws = wb.create_sheet(sn)
        ws.cell(1, 1, "时间\\到药材中心的距离")
        for j, d in enumerate(header):
            ws.cell(1, 2 + j, d)
        for i, t in enumerate(times):
            ws.cell(2 + i, 1, int(t) if float(t).is_integer() else round(float(t), 4))
            for j in range(len(header)):
                v = arr[i, j]
                if v is not None and not (isinstance(v, float) and np.isnan(v)):
                    ws.cell(2 + i, 2 + j, round(float(v), 4))
    wb.save(path)
    print("  saved", os.path.basename(path), flush=True)


def save_tab(name, times, prof, extra=None):
    rows = []
    for k, tt in enumerate(times):
        row = [tt] + list(prof[k])
        if extra is not None:
            row = row + [extra[k]]
        rows.append(row)
    arr = np.array(rows, float)
    np.save(os.path.join(OUTDIR, name + ".npy"), arr)
    return arr


# ============================================================
# 问题 1 : 预热平衡阶段 0~1800 s, 附录2, 固定半径
# ============================================================
print("=== 问题1 ===", flush=True)
t1e = np.arange(1, 1801)
sol1, _ = sv.solve_problem(2, 1800.0, moving=False, N=N, t_eval=t1e)
T1 = sv.extract_profile(sol1, N, 'T', DC)
C1 = sv.extract_profile(sol1, N, 'C', DC)
write_grid(os.path.join(OUTDIR, "result1.xlsx"), {"温度": T1, "水分浓度": C1}, DC, t1e)

tab_t = [100, 300, 600, 900, 1200, 1500, 1800]
idx1 = [int(np.argmin(abs(sol1.t - tt))) for tt in tab_t]
Ti = sv.extract_profile(sol1, N, 'T', TAB_D)[idx1]
Ci = sv.extract_profile(sol1, N, 'C', TAB_D)[idx1]
save_tab("tab1", [sol1.t[i] for i in idx1], Ti)
save_tab("tab2", [sol1.t[i] for i in idx1], Ci)

# ============================================================
# 问题 2 : 附录3, 3 h, 固定半径
# ============================================================
print("=== 问题2 ===", flush=True)
t2e = np.arange(1, 10801)
sol2, _ = sv.solve_problem(3, 10800.0, moving=False, N=N, t_eval=t2e)
T2 = sv.extract_profile(sol2, N, 'T', DC)
C2 = sv.extract_profile(sol2, N, 'C', DC)
write_grid(os.path.join(OUTDIR, "result2.xlsx"), {"温度": T2, "水分浓度": C2}, DC, t2e)

tab_h2 = [0.5, 1.0, 1.5, 2.0, 2.5, 3.0]
idx2 = [int(np.argmin(abs(sol2.t - h * 3600))) for h in tab_h2]
Ti2 = sv.extract_profile(sol2, N, 'T', TAB_D)[idx2]
Ci2 = sv.extract_profile(sol2, N, 'C', TAB_D)[idx2]
save_tab("tab3", [sol2.t[i] for i in idx2], Ti2)
save_tab("tab4", [sol2.t[i] for i in idx2], Ci2)

# ============================================================
# 问题 3 : 固定半径, 附录3, 烘干至 max C < 0.15
# ============================================================
print("=== 问题3 ===", flush=True)
t_end = 120 * 3600
t3e = np.arange(60, t_end + 1, 60)
sol3, _ = sv.solve_problem(3, t_end, moving=False, N=N, t_eval=t3e)
tc3, _ = sv.crossing_time(sol3, N)
print("  烘干结束: %.1f s = %.4f h = %.3f d" % (tc3, tc3 / 3600, tc3 / 86400), flush=True)

C3 = sv.extract_profile(sol3, N, 'C', DC)
n3 = int(np.searchsorted(sol3.t, tc3, side='right'))          # 末行 = floor(t_dry/60)*60
times3 = sol3.t[:n3]
write_grid(os.path.join(OUTDIR, "result3.xlsx"), {"Sheet1": C3[:n3]}, DC, times3)

tab5_t = list(np.arange(6, int(tc3 / 3600) + 1, 6) * 3600) + [tc3]
C5 = np.vstack([sv.extract_profile(sol3, N, 'C', TAB_D)[int(np.argmin(abs(sol3.t - tt)))]
                for tt in tab5_t[:-1]] + [sv.profile_at_time(sol3, N, 'C', tc3, TAB_D)])
save_tab("tab5", tab5_t, C5)
np.save(os.path.join(OUTDIR, "tend3.npy"), np.array([tc3]))

# ============================================================
# 问题 4 : 移动边界(半径收缩), 附录4
# ============================================================
print("=== 问题4 ===", flush=True)
t4e = np.arange(60, t_end + 1, 60)
sol4, _ = sv.solve_problem(4, t_end, moving=True, N=N, t_eval=t4e)
tc4, _ = sv.crossing_time(sol4, N)
print("  烘干结束: %.1f s = %.4f h = %.3f d" % (tc4, tc4 / 3600, tc4 / 86400), flush=True)

C4 = sv.extract_profile(sol4, N, 'C', DC[:, None] if False else DC)   # 物质坐标
Csurf4 = sv.surface_value(sol4, N, 'C')                                # xi=1 表面
n4 = int(np.searchsorted(sol4.t, tc4, side='right'))
times4 = sol4.t[:n4]
# result4: 0.0..1.9(物质坐标) + 药材表面
hdr4 = list(DC[:-1]) + ["药材表面"]
C4out = np.hstack([C4[:n4, :20], Csurf4[:n4].reshape(-1, 1)])
wb = openpyxl.Workbook(); wb.remove(wb.active)
ws = wb.create_sheet("Sheet1")
ws.cell(1, 1, "时间\\到药材中心的距离")
for j, d in enumerate(hdr4):
    ws.cell(1, 2 + j, d)
for i, t in enumerate(times4):
    ws.cell(2 + i, 1, int(t))
    for j in range(len(hdr4)):
        v = C4out[i, j]
        if not (isinstance(v, float) and np.isnan(v)):
            ws.cell(2 + i, 2 + j, round(float(v), 4))
wb.save(os.path.join(OUTDIR, "result4.xlsx"))
print("  saved result4.xlsx", flush=True)

tab6_t = list(np.arange(6, int(tc4 / 3600) + 1, 6) * 3600) + [tc4]
C6 = np.vstack([sv.extract_profile(sol4, N, 'C', TAB_D[:4])[int(np.argmin(abs(sol4.t - tt)))]
                for tt in tab6_t[:-1]] + [sv.profile_at_time(sol4, N, 'C', tc4, TAB_D[:4])])
Csurf6 = np.concatenate([Csurf4[[int(np.argmin(abs(sol4.t - tt))) for tt in tab6_t[:-1]]],
                         sv.profile_at_time(sol4, N, 'C', tc4, [2.0])])
Rnow = sv.Rfunc(np.array(tab6_t, float)) * 100.0        # 当时实际半径 cm
save_tab("tab6", tab6_t, C6, Csurf6)
np.save(os.path.join(OUTDIR, "tab6_R.npy"), Rnow)
np.save(os.path.join(OUTDIR, "tend4.npy"), np.array([tc4]))

print("\n全部完成。输出目录:", OUTDIR)
