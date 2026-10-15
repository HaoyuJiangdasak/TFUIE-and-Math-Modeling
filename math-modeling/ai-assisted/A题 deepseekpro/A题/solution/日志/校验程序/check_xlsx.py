# -*- coding: utf-8 -*-
"""check_xlsx.py —— 交付结果文件与附件 3 模板的合规性自检。

扫描 ../问题一~问题四/result?.xlsx，检查工作表名、A1 表头、距离列名、末列列名、
时间步长、4 位小数、无空格无负值、含水率沿半径单调不增，结果写入 ../日志/xlsx_check.txt。
用法： 在本文件夹（日志/校验程序）下执行  python check_xlsx.py
"""
import io
import os
import glob
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))        # solution/
LOGDIR = os.path.dirname(HERE)                       # solution/日志/

SPEC = [('问题一/result1.xlsx', 1, ['温度', '水分浓度'], False),
        ('问题二/result2.xlsx', 1, ['温度', '水分浓度'], False),
        ('问题三/result3.xlsx', 60, ['Sheet1'], True),
        ('问题四/result4.xlsx', 60, ['Sheet1'], True)]

log = io.open(os.path.join(LOGDIR, 'xlsx_check.txt'), 'w', encoding='utf-8')
allpass = True


def ok(cond, msg):
    global allpass
    log.write('  [%s] %s\n' % ('PASS' if cond else 'FAIL', msg))
    allpass &= bool(cond)


for rel, step, sheets, is_moist in SPEC:
    path = os.path.join(ROOT, rel.replace('/', os.sep))
    log.write('=' * 68 + '\n%s\n' % rel)
    if not os.path.exists(path):
        ok(False, '文件不存在（请先在对应文件夹里运行 prob 程序）')
        continue
    wb = openpyxl.load_workbook(path)
    ok(wb.sheetnames == sheets, '工作表名 %r（期望 %r）' % (wb.sheetnames, sheets))
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        hdr, data = rows[0], rows[1:]
        ts = [r[0] for r in data]
        vals = [v for r in data for v in r[1:]]
        dist = list(hdr[1:])
        log.write('  工作表 %s：%d 行 x %d 列\n' % (ws.title, len(data), len(dist)))
        ok(hdr[0] == '时间\\到药材中心的距离', 'A1 表头')
        ok(dist[:-1] == [round(0.1 * k, 4) for k in range(len(dist) - 1)], '距离列标题 0..1.9')
        last_ok = '药材表面' if rel.endswith('result4.xlsx') else round(0.1 * (len(dist) - 1), 4)
        ok(hdr[-1] == last_ok, '末列标题 = %r' % hdr[-1])
        ok(all(b - a == step for a, b in zip(ts, ts[1:])),
           '时间步长恒为 %d s（%d..%d）' % (step, ts[0], ts[-1]))
        ok(all(v is not None for v in vals), '数据区无空格')
        ok(min(vals) >= 0, '数据区无负值')
        ok(max(abs(v - round(v, 4)) for v in vals) == 0, '所有数值不超过 4 位小数')
        if is_moist or '水分' in ws.title:
            ok(all(all(r[j] >= r[j + 1] - 1e-12 for j in range(1, len(r) - 1)) for r in data),
               '含水率沿半径单调不增')
        if '温度' in ws.title:
            ok(27.9 <= min(vals) and max(vals) <= 50.5,
               '温度范围 [%.4f, %.4f] degC' % (min(vals), max(vals)))

log.write('\n总体：%s\n' % ('全部通过' if allpass else '存在未通过项'))
log.close()
print(open(os.path.join(LOGDIR, 'xlsx_check.txt'), encoding='utf-8').read())
