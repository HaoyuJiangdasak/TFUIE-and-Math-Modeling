from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parent
BOOK = ROOT / "output" / "第一问_残差_Bootstrap_Sobol结果.xlsx"


def main() -> None:
    wb = load_workbook(BOOK)
    if "分析说明" in wb.sheetnames:
        del wb["分析说明"]
    ws = wb.create_sheet("分析说明", 0)
    rows = [
        ("第一问：残差、Bootstrap与Sobol分析", ""),
        ("项目", "说明"),
        ("残差定义", "观测值－模型预测值"),
        ("边界Bootstrap", "1000次循环移动块残差Bootstrap，块长16"),
        ("Sobol样本", "N=2048，6个参数，共16384次模型运行"),
        ("Sobol估计量", "Jansen一阶效应S1与总效应ST"),
        ("Sobol区间", "对采样行进行1000次Bootstrap所得95%区间"),
        ("重要限制", "参数区间是工程探索假设；附件没有提供参数概率分布或标准差"),
        ("详细说明", "请查看文件夹中的00_使用说明.txt和01_结果摘要.txt"),
    ]
    for row in rows:
        ws.append(row)
    ws.merge_cells("A1:B1")
    ws["A1"].font = Font(name="Microsoft YaHei", size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="17365D")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    navy = PatternFill("solid", fgColor="1F4E78")
    light_blue = PatternFill("solid", fgColor="D9EAF7")
    assumption_fill = PatternFill("solid", fgColor="FFF2CC")
    thin = Side(style="thin", color="C8D1DC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for sheet in wb.worksheets:
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "A2"
        for row in sheet.iter_rows():
            for cell in row:
                cell.font = Font(name="Microsoft YaHei", size=10, color="000000")
                cell.alignment = Alignment(vertical="center", wrap_text=False)
                cell.border = border
        for cell in sheet[1]:
            cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
            cell.fill = navy
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.row_dimensions[1].height = 24
        if sheet.max_row > 1 and sheet.max_column > 1:
            sheet.auto_filter.ref = sheet.dimensions
        for column in range(1, sheet.max_column + 1):
            values = [str(sheet.cell(row, column).value or "") for row in range(1, min(sheet.max_row, 300) + 1)]
            width = min(max(max((len(v) for v in values), default=8) + 2, 11), 34)
            sheet.column_dimensions[get_column_letter(column)].width = width
        for row in range(2, sheet.max_row + 1):
            if row % 2 == 0:
                for cell in sheet[row]:
                    cell.fill = light_blue
            for cell in sheet[row]:
                if isinstance(cell.value, float):
                    cell.number_format = "0.000000E+00" if abs(cell.value) < 1e-4 and cell.value != 0 else "0.000000"

    note = wb["分析说明"]
    note.freeze_panes = "A3"
    note.auto_filter.ref = f"A2:B{note.max_row}"
    note.column_dimensions["A"].width = 20
    note.column_dimensions["B"].width = 78
    for cell in note[2]:
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
        cell.fill = navy
    note["A8"].fill = assumption_fill
    note["B8"].fill = assumption_fill
    note["A8"].font = Font(name="Microsoft YaHei", size=10, bold=True, color="9C0006")
    note["B8"].font = Font(name="Microsoft YaHei", size=10, bold=True, color="9C0006")

    if "Sobol参数区间" in wb.sheetnames:
        assumption = wb["Sobol参数区间"]
        for row in assumption.iter_rows(min_row=2):
            for cell in row:
                cell.fill = assumption_fill

    wb.save(BOOK)


if __name__ == "__main__":
    main()
