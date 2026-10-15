"""Create the paper-ready Table 1 and Table 2 workbook from the CN output."""
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


SCRIPT_DIR = Path(__file__).resolve().parent
MODEL_DIR = SCRIPT_DIR.parents[1]
RESULTS_DIR = MODEL_DIR / "output" / "results"
SOURCE = RESULTS_DIR / "result1_cn.xlsx"
OUTPUT = RESULTS_DIR / "表1表2_CN结果_修正版.xlsx"
TIMES = [100, 300, 600, 900, 1200, 1500, 1800]
RADII = [0, 0.5, 1, 1.5, 2]


def read_samples(sheet_name: str) -> list[list[float]]:
    workbook = load_workbook(SOURCE, data_only=True, read_only=True)
    sheet = workbook[sheet_name]
    rows = sheet.iter_rows(values_only=True)
    header = next(rows)
    radius_to_index = {
        float(value): index for index, value in enumerate(header[1:], start=1)
    }
    selected = {}
    for row in rows:
        time_value = int(row[0])
        if time_value in TIMES:
            selected[time_value] = [
                float(row[radius_to_index[radius]]) for radius in RADII
            ]
    values = [selected[time_value] for time_value in TIMES]
    workbook.close()
    return values


def add_table(workbook: Workbook, title: str, sheet_name: str, unit: str, values: list[list[float]]) -> None:
    sheet = workbook.create_sheet(sheet_name)
    sheet.merge_cells("A1:F1")
    sheet["A1"] = title
    sheet["A1"].font = Font(name="Arial", size=14, bold=True)
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet["A2"] = f"时间/s（数值单位：{unit}）"
    for column, radius in enumerate(RADII, start=2):
        sheet.cell(2, column, radius)

    for row, (time_value, row_values) in enumerate(zip(TIMES, values), start=3):
        sheet.cell(row, 1, time_value)
        for column, value in enumerate(row_values, start=2):
            sheet.cell(row, column, value)

    thin = Side(style="thin", color="808080")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for row in sheet.iter_rows(min_row=2, max_row=2 + len(TIMES), min_col=1, max_col=6):
        for cell in row:
            cell.font = Font(name="Arial", size=11, bold=cell.row == 2)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border
            if cell.row == 2:
                cell.fill = header_fill
            elif cell.column >= 2:
                cell.number_format = "0.0000"

    sheet["A2"].comment = Comment(
        "Source: q1_model/output/results/result1_cn.xlsx; fine-grid CN/FVM solution.",
        "Codex",
    )
    sheet.column_dimensions["A"].width = 28
    for column in "BCDEF":
        sheet.column_dimensions[column].width = 14
    sheet.freeze_panes = "B3"
    sheet.sheet_view.showGridLines = False


def main() -> None:
    temperature = read_samples("温度")
    moisture = read_samples("水分浓度")
    workbook = Workbook()
    workbook.remove(workbook.active)
    add_table(workbook, "表1  30分钟内药材的温度", "表1-温度", "°C", temperature)
    add_table(workbook, "表2  30分钟内药材的水分浓度", "表2-水分浓度", "kg/kg", moisture)
    workbook.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
