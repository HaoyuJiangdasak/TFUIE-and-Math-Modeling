# -*- coding: utf-8 -*-
"""参数敏感性分析 + 模型改进量化（用于报告第 8 节"模型评价与改进方向"）

输出格式（供 build_report.py 直接解析）：
    SENS|<名称>|<P3 h>|<P4 h>|<相对基准 %>
另外附带：仿射收缩体积–水量校核、问题1 补 Arrhenius 温度项对照。
"""
import numpy as np
import solver as sv
import time
import os

T_END = 120 * 3600
N = 321
TAB_D = [0, 0.5, 1.0, 1.5, 2.0]


def dry(mode, moving):
    sol, _ = sv.solve_problem(mode, T_END, moving=moving, N=N,
                              t_eval=np.arange(60, T_END + 1, 60),
                              rtol=1e-6, atol=1e-8, max_step=3600.0)
    tc, _ = sv.crossing_time(sol, N)
    return tc / 3600.0


def run(tag, base=None):
    print(">> " + tag, flush=True)
    t0 = time.time()
    d3 = dry(3, False)
    t1_ = time.time()
    d4 = dry(4, True)
    if base is None:
        r3 = r4 = 0.0
    else:
        r3 = 100.0 * (d3 - base[0]) / base[0]
        r4 = 100.0 * (d4 - base[1]) / base[1]
    print("SENS|%s|%.4f|%.4f|%+.2f%% / %+.2f%%" % (tag, d3, d4, r3, r4), flush=True)
    print("    (~%.0fs ; P3 %.0fs / P4 %.0fs)" % (time.time() - t0, t1_ - t0, time.time() - t1_), flush=True)
    return d3, d4


if __name__ == "__main__":
    print("=== 参数敏感性 (N=%d) ===" % N, flush=True)
    base = run("基准")

    sv.hm = 4e-7;  run("hm × 0.5", base);  sv.hm = 8e-7
    sv.hm = 1.6e-6; run("hm × 2",  base);  sv.hm = 8e-7
    sv.h = 50.0;    run("h × 2",    base);  sv.h = 25.0
    sv.C_HOLD = 0.02; run("Ca = 0.02", base); sv.C_HOLD = 0.05

    sv.T_HOLD = 45.0; run("Tf = 45 ℃", base); sv.T_HOLD = 50.0
    sv.T_HOLD = 55.0; run("Tf = 55 ℃", base); sv.T_HOLD = 50.0

    # D 下限截断（抑制 C→0 时 D 断崖式塌缩带来的刚性）
    _orig = sv.props

    def props_floored(mode, C, T_C):
        rho, cp, k, D = _orig(mode, C, T_C)
        return rho, cp, k, np.maximum(D, 1e-11)
    sv.props = props_floored
    run("Dmin = 1e-11", base)
    sv.props = _orig

    # ---- 仿射(均匀)收缩校核：体积减少 vs 排出水体积 ----
    print("\n=== 仿射(均匀)收缩校核 ===", flush=True)
    import openpyxl
    L, R0 = 0.25, 0.02
    # 终态含水率：用 result4.xlsx 末行做【体积加权】(圆柱权 ~ r dr)
    ws = openpyxl.load_workbook(
        os.path.join(sv.BASE, "结果", "result4.xlsx"), data_only=True)["Sheet1"]
    vals = [ws.cell(ws.max_row, c).value for c in range(2, ws.max_column + 1)]
    r20 = np.arange(0, 1.9001, 0.1)
    rr = np.append(r20, 2.0); CC = np.append(np.array(vals[:20], float), float(vals[20]))
    Cbar = float(np.trapezoid(CC * rr, rr) / np.trapezoid(rr, rr))
    # 初始湿基密度取【附录4】rho(C0)=760+90*C0（问题4 用的是附录4 物性）
    rho0 = 760.0 + 90.0 * 2.55
    V0 = np.pi * R0 ** 2 * L
    m_dry = rho0 * V0 / (1.0 + 2.55)
    dW = m_dry * (2.55 - Cbar)
    print("CHK|体积加权终态含水率|Cbar=%.4f|m_dry=%.5f kg|-" % (Cbar, m_dry), flush=True)
    for lab, Rf in [("烘干结束 51.09 h", 1.200e-2), ("附件2 末值 72 h", 1.198e-2)]:
        dV = V0 - np.pi * Rf ** 2 * L
        print("CHK|%s|R=%.4f cm|dV=%.4e m3|dW=%.5f kg|Vw=%.4e m3|ratio=%.3f"
              % (lab, Rf * 100, dV, dW, dW / 1000.0, dV / (dW / 1000.0)), flush=True)

    # ---- 问题1：附录2 的 D 是否补 Arrhenius 温度项 ----
    print("\n=== 问题1：附录2 的 D 补 Arrhenius 温度项 ===", flush=True)
    Ea_R, Tref = 3987.0, 333.15

    def props_arr(mode, C, T_C):
        rho, cp, k, D = _orig(mode, C, T_C)
        if mode == 2:
            Tk = np.asarray(T_C, float) + 273.15
            D = D * np.exp(-Ea_R * (1.0 / Tk - 1.0 / Tref))
        return rho, cp, k, D

    sol_a, _ = sv.solve_problem(2, 1800.0, N=N, t_eval=np.array([1800.0]))
    C_a = sv.extract_profile(sol_a, N, 'C', TAB_D)[0]
    sv.props = props_arr
    sol_b, _ = sv.solve_problem(2, 1800.0, N=N, t_eval=np.array([1800.0]))
    C_b = sv.extract_profile(sol_b, N, 'C', TAB_D)[0]
    sv.props = _orig
    print("ARR|原始 D(C)        |" + "|".join("%.4f" % v for v in C_a), flush=True)
    print("ARR|D(C,T) Arrhenius |" + "|".join("%.4f" % v for v in C_b), flush=True)
    rel = 100.0 * (C_b - C_a) / C_a
    print("ARR|相对差异 %|" + "|".join("%+.2f%%" % v for v in rel), flush=True)

    print("\ndone", flush=True)
