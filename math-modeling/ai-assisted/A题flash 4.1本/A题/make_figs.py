# -*- coding: utf-8 -*-
"""从 结果/result*.xlsx 生成报告插图 (PNG, 白底, 适配浅色主题)"""
import os
import numpy as np
import openpyxl
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "DejaVu Sans"]
plt.rcParams["axes.unicode_minus"] = False
plt.rcParams["figure.dpi"] = 130
plt.rcParams["savefig.bbox"] = "tight"

BASE = os.path.dirname(os.path.abspath(__file__))
RES = os.path.join(BASE, "结果")
FIG = os.path.join(RES, "figs")
os.makedirs(FIG, exist_ok=True)
C_G, C_B, C_K = "#0f6e4f", "#b45309", "#0b2a45"


def load(path, sheet=None):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[0]
    dist = np.array([float(v) if isinstance(v, (int, float)) else np.nan for v in hdr[1:]], float)
    t = np.array([float(r[0]) for r in rows[1:]], float)
    A = np.array([[np.nan if v is None else float(v) for v in r[1:]] for r in rows[1:]], float)
    wb.close()
    return t, dist, A


# ---- 图1: 问题1 水分浓度剖面 ----
fig, ax = plt.subplots(1, 2, figsize=(9.2, 3.4))
t1, d1, C1 = load(os.path.join(RES, "result1.xlsx"), "水分浓度")
t2, d2, C2 = load(os.path.join(RES, "result2.xlsx"), "水分浓度")
for tt, c in zip([300, 600, 900, 1200, 1800], ["#9ecfbb", "#5fae8b", "#2f8f68", "#0f6e4f", "#0b2a45"]):
    i = int(np.argmin(abs(t1 - tt)))
    ax[0].plot(d1, C1[i], color=c, lw=1.6, label="%d s" % int(t1[i]))
for tt, c in zip([1800, 3600, 5400, 7200, 10800], ["#f0c9a0", "#e0a463", "#c98436", "#a8681f", "#7a4a12"]):
    i = int(np.argmin(abs(t2 - tt)))
    ax[1].plot(d2, C2[i], color=c, lw=1.6, label="%.1f h" % (t2[i] / 3600))
for a, ti in zip(ax, ["问题1：预热平衡阶段（0–1800 s）", "问题2：全过程前 3 h"]):
    a.set_xlabel("到药材中心的距离 / cm"); a.set_ylabel("水分浓度 / (kg/kg)")
    a.set_title(ti, fontsize=10.5); a.grid(alpha=.25); a.legend(fontsize=7.5, ncol=2)
    a.set_facecolor("#ffffff")
fig.patch.set_facecolor("#ffffff")
fig.tight_layout(); fig.savefig(os.path.join(FIG, "fig1_profile.png")); plt.close(fig)

# ---- 图2: 问题3 中心/表面 含水率随时间 ----
t3, d3, C3 = load(os.path.join(RES, "result3.xlsx"))
tend3 = float(np.load(os.path.join(RES, "tend3.npy"))[0])
fig, ax = plt.subplots(figsize=(6.6, 3.6))
ax.plot(t3 / 3600, C3[:, 0], color=C_K, lw=1.8, label="中心 r=0")
ax.plot(t3 / 3600, C3[:, 10], color=C_G, lw=1.8, label="r=1.0 cm")
ax.plot(t3 / 3600, C3[:, -1], color=C_B, lw=1.8, label="表面 r=2.0 cm")
ax.axhline(0.15, color="#c0392b", ls="--", lw=1.1)
ax.axvline(tend3 / 3600, color="#c0392b", ls=":", lw=1.1)
ax.annotate("烘干结束 %.2f h" % (tend3 / 3600), xy=(tend3 / 3600, 1.6),
            xytext=(tend3 / 3600 - 20, 1.75), color="#c0392b", fontsize=9.5)
ax.text(1, 0.17, "0.15 kg/kg 判据", color="#c0392b", fontsize=8.5)
ax.set_xlabel("时间 / h"); ax.set_ylabel("水分浓度 / (kg/kg)")
ax.set_title("问题3：各位置水分浓度随时间变化（固定半径）", fontsize=10.5)
ax.grid(alpha=.25); ax.legend(); ax.set_facecolor("#ffffff")
fig.patch.set_facecolor("#ffffff"); fig.tight_layout()
fig.savefig(os.path.join(FIG, "fig2_p3.png")); plt.close(fig)

# ---- 图3: 问题4 物质坐标含水率 + 半径收缩 ----
t4, d4, C4 = load(os.path.join(RES, "result4.xlsx"))
R4 = np.load(os.path.join(RES, "tab6_R.npy"))
tend4 = float(np.load(os.path.join(RES, "tend4.npy"))[0])
fig, ax = plt.subplots(1, 2, figsize=(9.4, 3.6))
for k, c in zip([0, 5, 10, 19], ["#0b2a45", "#0f6e4f", "#b45309", "#c0392b"]):
    ax[0].plot(t4 / 3600, C4[:, k], lw=1.7, color=c, label="r0=%.1f cm" % d4[k])
ax[0].plot(t4 / 3600, C4[:, -1], lw=1.7, color="#7b1fa2", ls="--", label="药材表面")
ax[0].axhline(0.15, color="#c0392b", ls="--", lw=1.0)
ax[0].axvline(tend4 / 3600, color="#c0392b", ls=":", lw=1.0)
ax[0].set_xlabel("时间 / h"); ax[0].set_ylabel("水分浓度 / (kg/kg)")
ax[0].set_title("问题4：物质坐标各位置含水率（%.2f h 结束）" % (tend4 / 3600), fontsize=10.5)
ax[0].grid(alpha=.25); ax[0].legend(fontsize=8); ax[0].set_facecolor("#ffffff")

ttR = np.array([0, 1800, 3600, 5400, 7200, 9000, 10800, 14400, 21600, 28800, 43200, 86400, 129600, 172800, 183914])
import openpyxl as _ox
wb = _ox.load_workbook(os.path.join(BASE, "附件", "附件2.xlsx"), data_only=True)
ws = wb["Sheet1"]
tA = np.array([ws.cell(r, 1).value for r in range(2, ws.max_row + 1)], float)
RA = np.array([ws.cell(r, 2).value for r in range(2, ws.max_row + 1)], float)
ax[1].plot(tA / 3600, RA, color=C_G, lw=1.9, label="附件2 实测 R(t)")
ax[1].axhline(1.2, color="#c0392b", ls="--", lw=1.0)
ax[1].text(20, 1.23, "收缩终值 1.198 cm", color="#c0392b", fontsize=8.5)
ax[1].set_xlabel("时间 / h"); ax[1].set_ylabel("药材半径 / cm")
ax[1].set_title("药材半径随时间的收缩", fontsize=10.5)
ax[1].grid(alpha=.25); ax[1].legend(); ax[1].set_facecolor("#ffffff")
fig.patch.set_facecolor("#ffffff"); fig.tight_layout()
fig.savefig(os.path.join(FIG, "fig3_p4.png")); plt.close(fig)

# ---- 图4: 网格收敛 ----
fig, ax = plt.subplots(1, 2, figsize=(9.2, 3.3))
Ns = np.array([81, 161, 321, 641, 1281])
Cs = np.array([1.510378, 1.510270, 1.510244, 1.510237, 1.510235])
ax[0].semilogx(Ns, Cs, "o-", color=C_G, lw=1.7, ms=5)
ax[0].axhline(1.510234, color=C_B, ls="--", lw=1.0)
ax[0].text(90, 1.51030, "收敛值 1.51023", color=C_B, fontsize=9)
ax[0].set_xlabel("网格数 N"); ax[0].set_ylabel("C(r=2 cm, t=1800 s)")
ax[0].set_title("问题1 表面水分 网格收敛", fontsize=10.5)
ax[0].grid(alpha=.25, which="both"); ax[0].set_facecolor("#ffffff")
ax[1].semilogx([321, 641, 1281], [57.4936, 57.4788, 57.4742], "s-", color=C_K, lw=1.7, ms=5, label="问题3")
ax[1].semilogx([321, 641, 1281], [51.0900, 51.0878, 51.0873], "o-", color=C_B, lw=1.7, ms=5, label="问题4")
ax[1].set_xlabel("网格数 N"); ax[1].set_ylabel("烘干时长 / h")
ax[1].set_title("烘干时长 网格收敛", fontsize=10.5)
ax[1].grid(alpha=.25, which="both"); ax[1].legend(); ax[1].set_facecolor("#ffffff")
fig.patch.set_facecolor("#ffffff"); fig.tight_layout()
fig.savefig(os.path.join(FIG, "fig4_conv.png")); plt.close(fig)

print("figs done:", os.listdir(FIG))
