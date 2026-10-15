# -*- coding: utf-8 -*-
"""
生成 result1~result4.xlsx 及论文表1~表6
"""
import numpy as np
import openpyxl
import os
import solver as sv

BASE = sv.BASE
OUTDIR = os.path.join(BASE, "结果")
os.makedirs(OUTDIR, exist_ok=True)

N = 141
RDC = dict()  # 4位小数辅助
def r4(x):
    return np.round(x, 4)

# ---------------- 通用写表 ----------------
def write_grid(path, sheets, dists, times, minus1=None):
    """sheets: dict {sheet_name: 2D array (nt, nd)}, dists: 表头(nd), times: 时间列(nt)"""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    first = True
    for sn, arr in sheets.items():
        ws = wb.create_sheet(sn) if not first else wb.create_sheet(sn)
        first = False
        ws.cell(1, 1, "时间\\到药材中心的距离")
        for j, d in enumerate(dists):
            ws.cell(1, 2 + j, d)
        for i, t in enumerate(times):
            ws.cell(2 + i, 1, t)
            for j in range(len(dists)):
                v = arr[i, j]
                if not (v is None or (isinstance(v, float) and np.isnan(v))):
                    ws.cell(2 + i, 2 + j, round(float(v), 4))
    wb.save(path)
    print("saved", path)

# ============================================================
# 问题 1 : 预热阶段 0~1800s, 附录2
# ============================================================
print("=== 问题1 ===")
t1e = np.arange(1, 1801)
sol1, N1 = sv.solve_problem(2, 1800.0, moving=False, N=N, t_eval=t1e)
dc1 = np.round(np.arange(0, 2.0 + 1e-9, 0.1), 1)
T1 = sv.extract_profile(sol1, N, 'T', dc1)
C1 = sv.extract_profile(sol1, N, 'C', dc1)
write_grid(os.path.join(OUTDIR, "result1.xlsx"),
           {"温度": T1, "水分浓度": C1}, dc1, t1e)

# 表1 表2
tab_t = [100, 300, 600, 900, 1200, 1500, 1800]
tab_d = [0, 0.5, 1, 1.5, 2]
Ti = sv.extract_profile(sol1, N, 'T', tab_d)
Ci = sv.extract_profile(sol1, N, 'C', tab_d)
idx = [int(np.argmin(abs(sol1.t - tt))) for tt in tab_t]
tab1 = [[sol1.t[i]] + list(Ti[i]) for i in idx]
tab2 = [[sol1.t[i]] + list(Ci[i]) for i in idx]
np.save(os.path.join(OUTDIR, "tab1.npy"), np.array(tab1))
np.save(os.path.join(OUTDIR, "tab2.npy"), np.array(tab2))

# ============================================================
# 问题 2 : 附录3, 3h
# ============================================================
print("=== 问题2 ===")
t2e = np.arange(1, 10801)
sol2, N2 = sv.solve_problem(3, 10800.0, moving=False, N=N, t_eval=t2e)
T2 = sv.extract_profile(sol2, N, 'T', dc1)
C2 = sv.extract_profile(sol2, N, 'C', dc1)
write_grid(os.path.join(OUTDIR, "result2.xlsx"),
           {"温度": T2, "水分浓度": C2}, dc1, t2e)

tab_t2 = [0.5, 1.0, 1.5, 2.0, 2.5, 3.0]
tab_s2 = [int(h * 3600) for h in tab_t2]
Ti2 = sv.extract_profile(sol2, N, 'T', tab_d)
Ci2 = sv.extract_profile(sol2, N, 'C', tab_d)
idx2 = [int(np.argmin(abs(sol2.t - tt))) for tt in tab_s2]
tab3 = [[sol2.t[i]] + list(Ti2[i]) for i in idx2]
tab4 = [[sol2.t[i]] + list(Ci2[i]) for i in idx2]
np.save(os.path.join(OUTDIR, "tab3.npy"), np.array(tab3))
np.save(os.path.join(OUTDIR, "tab4.npy"), np.array(tab4))

# ============================================================
# 问题 3 : 烘干至 max C < 0.15, 固定半径, 附录3
# ============================================================
print("=== 问题3 ===")
t_end = 120 * 3600
t3e = np.arange(60, t_end + 1, 60)
sol3, N3 = sv.solve_problem(3, t_end, moving=False, N=N, t_eval=t3e)
Cmax3 = sol3.y[N:, :].max(axis=0)
below = np.where(Cmax3 < 0.15)[0]
tend3_idx = int(below[0])
tend3 = sol3.t[tend3_idx]
print("问题3 烘干结束:", round(tend3 / 3600, 3), "h")

C3 = sv.extract_profile(sol3, N, 'C', dc1)
# 截断到烘干结束(含)
times3 = sol3.t[:tend3_idx + 1]
C3c = C3[:tend3_idx + 1]
write_grid(os.path.join(OUTDIR, "result3.xlsx"), {"Sheet1": C3c}, dc1, times3)

# 表5: 每隔6h + 烘干结束
tab5_times = list(np.arange(6, int(tend3 / 3600) + 1, 6) * 3600)
tab5_times.append(tend3)
C5 = sv.extract_profile(sol3, N, 'C', tab_d)
tab5 = []
for tt in tab5_times:
    i = int(np.argmin(abs(sol3.t - tt)))
    tab5.append([sol3.t[i]] + list(C5[i]))
np.save(os.path.join(OUTDIR, "tab5.npy"), np.array(tab5))
np.save(os.path.join(OUTDIR, "tend3.npy"), np.array([tend3]))

# ============================================================
# 问题 4 : 移动边界(半径收缩), 附录4
# ============================================================
print("=== 问题4 ===")
t4e = np.arange(60, t_end + 1, 60)
sol4, N4 = sv.solve_problem(4, t_end, moving=True, N=N, t_eval=t4e)
Cmax4 = sol4.y[N:, :].max(axis=0)
below4 = np.where(Cmax4 < 0.15)[0]
tend4_idx = int(below4[0])
tend4 = sol4.t[tend4_idx]
print("问题4 烘干结束:", round(tend4 / 3600, 3), "h")

# result4: 固定距离 0..1.9 + 药材表面列
dcfix = np.round(np.arange(0, 1.9 + 1e-9, 0.1), 1)   # 0..1.9 共20列
times4 = sol4.t[:tend4_idx + 1]
Cfix_full = sv.extract_profile(sol4, N, 'C', dcfix, moving=True)  # 空白处为 NaN
Cfix = Cfix_full[:tend4_idx + 1]
Csurf = sol4.y[N + (N - 1), :tend4_idx + 1]           # 表面值 (= r=R(t))
# 组装 数组 (nt, 21): 前20列固定距离, 最后1列药材表面
C4out = np.hstack([Cfix, Csurf.reshape(-1, 1)])
d4header = list(dcfix) + ["药材表面"]
# 写 result4
wb = openpyxl.Workbook(); wb.remove(wb.active)
ws = wb.create_sheet("Sheet1")
ws.cell(1, 1, "时间\\到药材中心的距离")
for j, d in enumerate(d4header):
    ws.cell(1, 2 + j, d)
for i, t in enumerate(times4):
    ws.cell(2 + i, 1, t)
    for j in range(len(d4header)):
        v = C4out[i, j]
        if not (isinstance(v, float) and np.isnan(v)):
            ws.cell(2 + i, 2 + j, round(float(v), 4))
wb.save(os.path.join(OUTDIR, "result4.xlsx"))
print("saved", os.path.join(OUTDIR, "result4.xlsx"))

# 表6: 每隔6h + 烘干结束, 距离 0,0.5,1,1.5 + 药材表面
tab6_times = list(np.arange(6, int(tend4 / 3600) + 1, 6) * 3600)
tab6_times.append(tend4)
tab6_d = [0, 0.5, 1.0, 1.5]
C6 = sv.extract_profile(sol4, N, 'C', tab6_d, moving=True)
Csurf_all = sol4.y[N + (N - 1), :]
tab6 = []
for tt in tab6_times:
    i = int(np.argmin(abs(sol4.t - tt)))
    row = [sol4.t[i]] + list(np.where(np.isnan(C6[i]), np.nan, C6[i])) + [Csurf_all[i]]
    tab6.append(row)
np.save(os.path.join(OUTDIR, "tab6.npy"), np.array(tab6, dtype=float))
np.save(os.path.join(OUTDIR, "tend4.npy"), np.array([tend4]))

print("\n全部完成。输出目录:", OUTDIR)