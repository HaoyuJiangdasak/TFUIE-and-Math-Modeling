# -*- coding: utf-8 -*-
"""
2026 CUMCM A题《药材的烘干问题》数值求解器
圆柱(径向一维) 热传导 + 水分扩散 耦合模型, 有限体积法(ξ=r/R 归一化坐标)
支持时变边界条件(附件1) 与 移动边界(附件2 半径收缩)
"""
import numpy as np
from scipy.integrate import solve_ivp
import openpyxl
import os

# ---------------- 常数 ----------------
R0 = 0.02        # 初始半径 m
T0 = 28.0        # 初始温度 °C
C0 = 2.55        # 初始干基含水率 kg/kg

h  = 25.0        # 对流换热系数 W/(m^2 K)
hm = 8e-7        # 对流传质系数 m/s

# ---------------- 物性参数 ----------------
def props(mode, C, T_C):
    """返回 rho, cp, k, D"""
    Tk = T_C + 273.15
    if mode == 2:          # 附录2 (问题1)
        rho = 820.0
        cp  = 2600.0
        k   = 0.36
        D   = 7e-9 * np.exp(-0.89 / C)
    elif mode == 3:        # 附录3 (问题2,3)
        rho = 650.0 + 128.0 * C
        cp  = 1450.0 + 2736.0 * C / (C + 1.0)
        k   = 0.21 + 0.38 * C / (C + 1.0)
        D   = 2.4e-3 * np.exp(-0.45 / C) * np.exp(-3850.0 / Tk)
    elif mode == 4:        # 附录4 (问题4)
        rho = 760.0 + 90.0 * C
        cp  = 1850.0 + 2150.0 * C / (C + 1.0)
        k   = 0.12 + 0.20 * C / (C + 1.0)
        D   = 4.2e-4 * np.exp(-0.30 / C) * np.exp(-3850.0 / Tk)
    else:
        raise ValueError("unknown mode")
    return rho, cp, k, D


# ---------------- 附件数据 ----------------
BASE = os.path.dirname(os.path.abspath(__file__))
ATT  = os.path.join(BASE, "附件")

_wb1 = openpyxl.load_workbook(os.path.join(ATT, "附件1.xlsx"), data_only=True)
_ws1 = _wb1["Sheet1"]
t1, Ta, Ca = [], [], []
for r in range(2, _ws1.max_row + 1):
    t1.append(_ws1.cell(r, 1).value)
    Ta.append(_ws1.cell(r, 2).value)
    Ca.append(_ws1.cell(r, 3).value)
t1 = np.array(t1, float); Ta = np.array(Ta, float); Ca = np.array(Ca, float)

_wb2 = openpyxl.load_workbook(os.path.join(ATT, "附件2.xlsx"), data_only=True)
_ws2 = _wb2["Sheet1"]
t2, Rad = [], []
for r in range(2, _ws2.max_row + 1):
    t2.append(_ws2.cell(r, 1).value)
    Rad.append(_ws2.cell(r, 2).value)
t2 = np.array(t2, float); Rad = np.array(Rad, float)   # Rad 单位 cm

def Tair(t):
    """烘房温度 °C: 0~14400s 用附件1插值, 之后恒温 50°C"""
    t = np.asarray(t, float)
    return np.where(t <= 14400.0, np.interp(t, t1, Ta), 50.0)

def Cair(t):
    """烘房水分浓度 kg/kg: 0~14400s 附件1插值, 之后恒 0.05"""
    t = np.asarray(t, float)
    return np.where(t <= 14400.0, np.interp(t, t1, Ca), 0.05)

def Rfunc(t):
    """半径(米), 附件2插值, 超出取末值"""
    t = np.asarray(t, float)
    v = np.interp(np.clip(t, 0.0, t2[-1]), t2, Rad)
    return v / 100.0

def dRdt(t):
    """dR/dt (m/s), 附件2分段差分, 末尾取0"""
    t = np.asarray(t, float)
    out = []
    for tt in np.atleast_1d(t):
        if tt >= t2[-1]:
            out.append(0.0); continue
        i = np.searchsorted(t2, tt)
        i = max(1, min(i, len(t2) - 1))
        dt_ = t2[i] - t2[i - 1]
        out.append((Rad[i] - Rad[i - 1]) / dt_ * 1e-2)
    out = np.array(out)
    return out if out.size > 1 else float(out[0])


# ---------------- 统一 FVM 离散 (xi = r/R) ----------------
def make_rhs(N, mode, moving=False):
    dxi = 1.0 / (N - 1)
    xi = np.linspace(0.0, 1.0, N)          # 节点 xi
    xif = (np.arange(N - 1) + 0.5) * dxi   # 内部面 xi_{i+1/2}, i=0..N-2

    def rhs(t, y):
        T = y[:N]; C = y[N:]
        R = Rfunc(t) if moving else R0
        dR = dRdt(t) if moving else 0.0
        Ta_ = float(Tair(t) if np.ndim(Tair(t)) else Tair(t))
        Ca_ = float(Cair(t))

        rho = np.empty(N); cp = np.empty(N); kk = np.empty(N); D = np.empty(N)
        for i in range(N):
            rho[i], cp[i], kk[i], D[i] = props(mode, C[i], T[i])

        # 内部面物性(平均)
        Cf = 0.5 * (C[:-1] + C[1:]); Tf = 0.5 * (T[:-1] + T[1:])
        kf = np.empty(N - 1); Df = np.empty(N - 1)
        for i in range(N - 1):
            _, _, kf[i], Df[i] = props(mode, Cf[i], Tf[i])

        dT = np.empty(N); dC = np.empty(N)
        invR2 = 1.0 / (R * R)

        # ---- 中心节点 i=0 ----
        dT[0] = 4.0 * kk[0] / (rho[0] * cp[0]) * invR2 / (dxi * dxi) * (T[1] - T[0])
        dC[0] = 4.0 * D[0] * invR2 / (dxi * dxi) * (C[1] - C[0])

        # ---- 内部节点 i=1..N-2 ----
        for i in range(1, N - 1):
            # 西面 i-1/2, 东面 i+1/2
            fx_e = xif[i] * kf[i] * (T[i + 1] - T[i]) / dxi
            fx_w = xif[i - 1] * kf[i - 1] * (T[i] - T[i - 1]) / dxi
            dT[i] = (1.0 / xi[i]) * (fx_e - fx_w) / dxi * invR2 / (rho[i] * cp[i])
            dT[i] += (xi[i] * dR / R) * (T[i + 1] - T[i - 1]) / (2 * dxi)

            fc_e = xif[i] * Df[i] * (C[i + 1] - C[i]) / dxi
            fc_w = xif[i - 1] * Df[i - 1] * (C[i] - C[i - 1]) / dxi
            dC[i] = (1.0 / xi[i]) * (fc_e - fc_w) / dxi * invR2
            dC[i] += (xi[i] * dR / R) * (C[i + 1] - C[i - 1]) / (2 * dxi)

        # ---- 表面节点 i=N-1 (ghost + Robin) ----
        i = N - 1
        # ghost 值 (在 xi = 1 + dxi)
        T_g = T[i - 1] - (2 * dxi * R * h / kk[i]) * (T[i] - Ta_)
        C_g = C[i - 1] - (2 * dxi * R * hm / D[i]) * (C[i] - Ca_)
        # 西面 (i-1/2) 与 东面 (i+1/2, 用ghost)
        xi_e = 1.0 + 0.5 * dxi
        xi_w = 1.0 - 0.5 * dxi
        kf_e = 0.5 * (kk[i] + kk[i - 1]); kf_w = kf[i - 1]
        Df_e = 0.5 * (D[i] + D[i - 1]); Df_w = Df[i - 1]
        fx_e = xi_e * kf_e * (T_g - T[i]) / dxi
        fx_w = xi_w * kf_w * (T[i] - T[i - 1]) / dxi
        dT[i] = (1.0 / xi[i]) * (fx_e - fx_w) / dxi * invR2 / (rho[i] * cp[i])
        dT[i] += (xi[i] * dR / R) * (T_g - T[i - 1]) / (2 * dxi)

        fc_e = xi_e * Df_e * (C_g - C[i]) / dxi
        fc_w = xi_w * Df_w * (C[i] - C[i - 1]) / dxi
        dC[i] = (1.0 / xi[i]) * (fc_e - fc_w) / dxi * invR2
        dC[i] += (xi[i] * dR / R) * (C_g - C[i - 1]) / (2 * dxi)

        return np.concatenate([dT, dC])

    return rhs


def solve_problem(mode, t_end, moving=False, N=161, t_eval=None, rtol=1e-7, atol=1e-9):
    rhs = make_rhs(N, mode, moving)
    y0 = np.concatenate([np.full(N, T0), np.full(N, C0)])
    if t_eval is None:
        t_eval = np.linspace(0, t_end, 2)
    sol = solve_ivp(rhs, (0.0, t_end), y0, method='BDF', t_eval=t_eval,
                    rtol=rtol, atol=atol, max_step=np.inf)
    return sol, N


def extract_profile(sol, N, field, r_cm, moving=False):
    """提取 field('T'或'C') 在指定半径(厘米)处的值 -> (nt, nr) 数组"""
    xi = np.linspace(0, 1, N)
    t = np.asarray(sol.t, float)
    rq = np.asarray(r_cm, float)
    nt = len(t); nr = len(rq)
    out = np.full((nt, nr), np.nan)
    Rt = Rfunc(t) if moving else np.full(nt, R0)
    for j in range(nt):
        r_phys = xi * Rt[j] * 100.0          # cm, 单调 0..R
        yv = sol.y[:N, j] if field == 'T' else sol.y[N:, j]
        if moving:
            mask = rq <= r_phys[-1] + 1e-12
            out[j, mask] = np.interp(rq[mask], r_phys, yv)
        else:
            out[j] = np.interp(rq, r_phys, yv)
    return out