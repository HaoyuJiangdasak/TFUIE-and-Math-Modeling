from __future__ import annotations

from dataclasses import dataclass, field, replace
from pathlib import Path
from typing import Literal, Optional, Sequence
import argparse
import math

import numpy as np
import pandas as pd
from scipy.integrate import solve_ivp, cumulative_trapezoid
from scipy.optimize import brentq
from openpyxl import Workbook, load_workbook

from q2_solver import (
    Array,
    FaceAverage,
    Q2Parameters,
    SolverConfig,
    PicardConfig,
    FVMGrid,
    EnvironmentData,
    density,
    heat_capacity,
    thermal_conductivity,
    moisture_diffusivity,
    check_material_state,
    coupled_surface_values,
    joint_rhs,
    jacobian_sparsity,
    center_reconstruction,
    sample_state_pair,
    _be_diffusion_solve,
    _k_faces,
    _D_faces,
)

BoundaryStrategy = Literal["last_hour_mean", "last_value", "last_1p5h_mean"]


@dataclass(frozen=True)
class Q3Parameters(Q2Parameters):
    """Q3 uses the same fixed-domain material/transfer parameters as Q2."""
    threshold: float = 0.15


@dataclass(frozen=True)
class Q3SolverConfig:
    """Long-time/event configuration layered on the validated Q2 spatial core."""
    method: str = "BDF"
    rtol: float = 1.0e-9
    atol: float = 1.0e-11

    # The measured 0-4 h boundary still has 60 s knots, so retain the Q2 step cap.
    max_step_data: float = 5.0

    # After 4 h the boundary is constant and dynamics are much slower.
    max_step_long: float = 60.0

    # result3.xlsx output cadence; this is NOT the solver time step.
    output_dt: float = 60.0
    output_dr: float = 0.001  # 0.1 cm

    face_average: FaceAverage = "harmonic"
    segment_at_data_knots: bool = True

    # Long-time chunks keep memory bounded and provide full-state monitoring.
    long_segment_s: float = 3600.0

    # Local coupled Robin Picard controls inherited in spirit from Q2.
    surface_tol_T: float = 1.0e-11
    surface_tol_C: float = 1.0e-12
    surface_max_iter: int = 50
    surface_relaxation: float = 1.0

    # Terminal-event refinement.
    refine_event: bool = True
    event_rtol: float = 2.0e-10
    event_atol: float = 2.0e-12
    event_max_step: float = 5.0
    event_refine_window_s: float = 300.0

    # Safety horizon. Raise explicitly if a chosen model has not dried by then.
    max_time_s: float = 15.0 * 24.0 * 3600.0  # 15 days

    def rhs_config(self) -> SolverConfig:
        # max_step/output_dt are supplied by the Q3 driver, not joint_rhs itself.
        return SolverConfig(
            method=self.method,
            rtol=self.rtol,
            atol=self.atol,
            max_step=self.max_step_data,
            output_dt=self.output_dt,
            face_average=self.face_average,
            segment_at_data_knots=self.segment_at_data_knots,
            surface_tol_T=self.surface_tol_T,
            surface_tol_C=self.surface_tol_C,
            surface_max_iter=self.surface_max_iter,
            surface_relaxation=self.surface_relaxation,
        )


@dataclass(frozen=True)
class PlateauStatistics:
    strategy: BoundaryStrategy
    data_end_s: float
    window_start_s: float
    window_end_s: float
    n_points: int
    T_long_C: float
    C_long: float
    T_std_C: float
    C_std: float
    T_slope_C_per_h: float
    C_slope_per_h: float


@dataclass(frozen=True)
class LongTermEnvironment:
    """Measured boundary on [0, data_end], then a constant plateau extension."""
    base: EnvironmentData
    strategy: BoundaryStrategy
    stats: PlateauStatistics

    @classmethod
    def from_excel(
        cls,
        path: str | Path,
        strategy: BoundaryStrategy = "last_hour_mean",
    ) -> "LongTermEnvironment":
        base = EnvironmentData.from_excel(path)
        t_end = float(base.t[-1])

        if strategy == "last_hour_mean":
            start = t_end - 3600.0
        elif strategy == "last_value":
            start = t_end
        elif strategy == "last_1p5h_mean":
            start = t_end - 5400.0
        else:
            raise ValueError(f"Unknown long-term boundary strategy: {strategy}")

        if strategy == "last_value":
            mask = np.isclose(base.t, t_end, atol=1e-10, rtol=0)
        else:
            mask = (base.t >= start - 1e-12) & (base.t <= t_end + 1e-12)

        tw = base.t[mask]
        Tw = base.temperature[mask]
        Cw = base.moisture[mask]
        if tw.size == 0:
            raise ValueError("Selected plateau window contains no measured data.")

        if strategy == "last_value":
            T_long = float(Tw[-1])
            C_long = float(Cw[-1])
            T_std = 0.0
            C_std = 0.0
            T_slope = 0.0
            C_slope = 0.0
        else:
            T_long = float(np.mean(Tw))
            C_long = float(np.mean(Cw))
            T_std = float(np.std(Tw, ddof=1)) if tw.size > 1 else 0.0
            C_std = float(np.std(Cw, ddof=1)) if tw.size > 1 else 0.0
            xh = tw / 3600.0
            T_slope = float(np.polyfit(xh, Tw, 1)[0]) if tw.size > 1 else 0.0
            C_slope = float(np.polyfit(xh, Cw, 1)[0]) if tw.size > 1 else 0.0

        stats = PlateauStatistics(
            strategy=strategy,
            data_end_s=t_end,
            window_start_s=float(start),
            window_end_s=t_end,
            n_points=int(tw.size),
            T_long_C=T_long,
            C_long=C_long,
            T_std_C=T_std,
            C_std=C_std,
            T_slope_C_per_h=T_slope,
            C_slope_per_h=C_slope,
        )
        return cls(base=base, strategy=strategy, stats=stats)

    @property
    def t(self) -> Array:
        # Compatibility with Q2 helpers that inspect measured boundary knots.
        return self.base.t

    @property
    def temperature(self) -> Array:
        return self.base.temperature

    @property
    def moisture(self) -> Array:
        return self.base.moisture

    @property
    def data_end_s(self) -> float:
        return self.stats.data_end_s

    def T_inf(self, t: float | Array) -> Array:
        x = np.asarray(t, dtype=float)
        measured = np.interp(np.minimum(x, self.data_end_s), self.base.t, self.base.temperature)
        out = np.where(x <= self.data_end_s + 1e-12, measured, self.stats.T_long_C)
        return out

    def C_inf(self, t: float | Array) -> Array:
        x = np.asarray(t, dtype=float)
        measured = np.interp(np.minimum(x, self.data_end_s), self.base.t, self.base.moisture)
        out = np.where(x <= self.data_end_s + 1e-12, measured, self.stats.C_long)
        return out


@dataclass
class Q3Solution:
    grid: FVMGrid
    env: LongTermEnvironment
    params: Q3Parameters
    config: Q3SolverConfig

    t: Array                  # regular result3 times: 0,60,...,floor(tdry/60)*60
    r: Array                  # 0,0.1,...,2.0 cm
    T: Array                  # sampled field, retained for analysis
    C: Array                  # sampled field, written to result3.xlsx

    # Sparse full-state monitoring snapshots.
    state_t: Array
    T_state: Array
    C_state: Array

    # Exact terminal event.
    t_dry_s: float
    T_dry_state: Array
    C_dry_state: Array
    Cmax_dry: float
    rmax_dry_m: float

    solver_stats: dict = field(default_factory=dict)


@dataclass(frozen=True)
class DryingTimeResult:
    N: int
    t_dry_s: float
    t_dry_h: float
    Cmax: float
    rmax_m: float
    solver_stats: dict


@dataclass(frozen=True)
class BlockPicardEventResult:
    N: int
    dt_s: float
    t_dry_s: float
    t_dry_h: float
    max_picard_iterations: int
    mean_picard_iterations: float



# -----------------------------------------------------------------------------
# Q3 low-moisture robust surface solver
# -----------------------------------------------------------------------------

def _surface_picard_map(
    Cs: float,
    Tp: float,
    Tpp: float,
    Cp: float,
    Cpp: float,
    Tinf: float,
    Cinf: float,
    dr: float,
    params: Q3Parameters,
) -> tuple[float, float]:
    """One local Gauss-Seidel Picard map Cs -> Ts -> Cs_raw."""
    ks = float(thermal_conductivity(Cs))
    Ts = (ks * (9.0 * Tp - Tpp) + 3.0 * dr * params.h * Tinf) / (
        8.0 * ks + 3.0 * dr * params.h
    )
    Ds = float(moisture_diffusivity(Cs, Ts))
    Cs_raw = (Ds * (9.0 * Cp - Cpp) + 3.0 * dr * params.hm * Cinf) / (
        8.0 * Ds + 3.0 * dr * params.hm
    )
    return float(Ts), float(Cs_raw)


def coupled_surface_values(
    T_C: Array,
    C: Array,
    t: float,
    grid: FVMGrid,
    env: LongTermEnvironment,
    params: Q3Parameters,
    config: SolverConfig,
    counter: Optional[dict] = None,
) -> tuple[float, float, int]:
    """Q3 local coupled Robin solve.

    Primary algorithm: the same local coupled Picard iteration chosen for Q2/Q3.
    At low moisture the scalar fixed-point map can cease to be contractive.
    If Picard does not converge, use a bracketed Brent solve of the *same*
    scalar fixed-point equation as a transparent safety fallback.

    The fallback changes only the local nonlinear equation solver; it does not
    change the PDE, FVM flux, Robin law, or time integrator.
    """
    if T_C.size < 2 or C.size < 2:
        raise ValueError("At least two cells are required.")

    Tp, Tpp = float(T_C[-1]), float(T_C[-2])
    Cp, Cpp = float(C[-1]), float(C[-2])
    Tinf = float(env.T_inf(t))
    Cinf = float(env.C_inf(t))
    dr = grid.dr

    omega = float(config.surface_relaxation)
    if not (0.0 < omega <= 1.0):
        raise ValueError("surface_relaxation must lie in (0,1].")

    # Physical branch continuation starts from the last interior value.
    Cs = max(Cp, 1e-12)
    Ts, _ = _surface_picard_map(Cs, Tp, Tpp, Cp, Cpp, Tinf, Cinf, dr, params)

    prev_delta = None
    dynamic_omega = omega
    for it in range(1, config.surface_max_iter + 1):
        Ts_raw, Cs_raw = _surface_picard_map(
            Cs, Tp, Tpp, Cp, Cpp, Tinf, Cinf, dr, params
        )
        delta = Cs_raw - Cs

        # If the fixed-point updates start alternating without shrinking,
        # damp them. This remains a Picard iteration and is cheap.
        if prev_delta is not None and delta * prev_delta < 0.0:
            if abs(delta) > 0.8 * abs(prev_delta):
                dynamic_omega = max(0.15, 0.5 * dynamic_omega)

        Cs_new = dynamic_omega * Cs_raw + (1.0 - dynamic_omega) * Cs
        Ts_new, _ = _surface_picard_map(
            Cs_new, Tp, Tpp, Cp, Cpp, Tinf, Cinf, dr, params
        )

        if Cs_new <= 0.0 or not (math.isfinite(Ts_new) and math.isfinite(Cs_new)):
            raise FloatingPointError(
                f"Surface Picard produced an invalid state at t={t}: "
                f"Ts={Ts_new}, Cs={Cs_new}"
            )

        if abs(Ts_new - Ts) <= config.surface_tol_T and abs(Cs_new - Cs) <= config.surface_tol_C:
            if counter is not None:
                counter["picard_success"] = counter.get("picard_success", 0) + 1
            return float(Ts_new), float(Cs_new), it

        prev_delta = delta
        Ts, Cs = Ts_new, Cs_new

    # Picard fallback: solve H(Cs)=Cs-F(Cs)=0 on the physical drying branch.
    # For the Q3 drying regime Cinf < Cs <= near-surface interior moisture.
    def residual(x: float) -> float:
        _, fx = _surface_picard_map(
            x, Tp, Tpp, Cp, Cpp, Tinf, Cinf, dr, params
        )
        return x - fx

    lower = max(min(Cinf, Cp, Cpp), 1e-12)
    upper = max(Cp, Cpp, Cinf) * (1.0 + 1e-10)
    if upper <= lower:
        upper = lower + max(1e-8, abs(lower) * 1e-6)

    # The nonlinear Robin equation can possess several mathematical roots
    # (this is already visible at t=0). Picard started from Cp follows the
    # high-moisture root continuously. Therefore the fallback must NOT call
    # Brent on the entire interval blindly; it scans for all sign changes and
    # selects the largest root, i.e. the branch connected to the near-surface
    # interior state.
    scan = np.linspace(lower, upper, 96)
    fscan = np.array([residual(float(x)) for x in scan])
    brackets: list[tuple[float, float]] = []
    for jj in range(len(scan) - 1):
        if fscan[jj] == 0.0:
            brackets.append((float(scan[jj]), float(scan[jj])))
        elif fscan[jj] * fscan[jj + 1] < 0.0:
            brackets.append((float(scan[jj]), float(scan[jj + 1])))

    if fscan[-1] == 0.0:
        brackets.append((float(scan[-1]), float(scan[-1])))

    if not brackets:
        # Retry with a denser scan before declaring a local boundary failure.
        scan = np.linspace(lower, upper, 512)
        fscan = np.array([residual(float(x)) for x in scan])
        for jj in range(len(scan) - 1):
            if fscan[jj] == 0.0:
                brackets.append((float(scan[jj]), float(scan[jj])))
            elif fscan[jj] * fscan[jj + 1] < 0.0:
                brackets.append((float(scan[jj]), float(scan[jj + 1])))

    if not brackets:
        raise RuntimeError(
            f"Surface Picard failed and no physical fixed-point bracket was found "
            f"at t={t:.6g} s; Cp={Cp:.6g}, Cpp={Cpp:.6g}, Cinf={Cinf:.6g}."
        )

    a, b = brackets[-1]  # largest/continuation root
    if a == b:
        root = a
    else:
        root = float(brentq(
            residual, a, b,
            xtol=max(config.surface_tol_C, 1e-14),
            rtol=1e-12, maxiter=100,
        ))

    Ts_root, _ = _surface_picard_map(
        root, Tp, Tpp, Cp, Cpp, Tinf, Cinf, dr, params
    )
    if counter is not None:
        counter["brent_fallback"] = counter.get("brent_fallback", 0) + 1
    return float(Ts_root), float(root), config.surface_max_iter


def joint_rhs(
    t: float,
    y: Array,
    grid: FVMGrid,
    env: LongTermEnvironment,
    params: Q3Parameters,
    config: SolverConfig,
    surface_counter: Optional[dict] = None,
) -> Array:
    """Q2 semi-discrete equations with the Q3 robust long-time surface solve."""
    N = grid.N
    T_C = np.asarray(y[:N], dtype=float)
    C = np.asarray(y[N:], dtype=float)
    check_material_state(T_C, C)

    kf = _k_faces(C, config.face_average)
    Df = _D_faces(C, T_C, config.face_average)

    q = np.zeros(N + 1, dtype=float)
    J = np.zeros(N + 1, dtype=float)
    q[1:N] = -kf * (T_C[1:] - T_C[:-1]) / grid.dr
    J[1:N] = -Df * (C[1:] - C[:-1]) / grid.dr

    local_counter: dict = {}
    Ts, Cs, niter = coupled_surface_values(
        T_C, C, t, grid, env, params, config, counter=local_counter
    )
    q[N] = params.h * (Ts - float(env.T_inf(t)))
    J[N] = params.hm * (Cs - float(env.C_inf(t)))

    if surface_counter is not None:
        surface_counter["calls"] = surface_counter.get("calls", 0) + 1
        surface_counter["iterations"] = surface_counter.get("iterations", 0) + niter
        surface_counter["max_iterations"] = max(
            surface_counter.get("max_iterations", 0), niter
        )
        surface_counter["picard_success"] = (
            surface_counter.get("picard_success", 0)
            + local_counter.get("picard_success", 0)
        )
        surface_counter["brent_fallback"] = (
            surface_counter.get("brent_fallback", 0)
            + local_counter.get("brent_fallback", 0)
        )

    net_q = grid.areas[:-1] * q[:-1] - grid.areas[1:] * q[1:]
    net_J = grid.areas[:-1] * J[:-1] - grid.areas[1:] * J[1:]
    cap = density(C) * heat_capacity(C)

    dTdt = net_q / (cap * grid.volumes)
    dCdt = net_J / grid.volumes
    return np.concatenate([dTdt, dCdt])


def sample_state_pair(
    T_C: Array,
    C: Array,
    t: float,
    target_r: Array,
    grid: FVMGrid,
    env: LongTermEnvironment,
    params: Q3Parameters,
    config: SolverConfig,
) -> tuple[Array, Array]:
    target_r = np.asarray(target_r, dtype=float)
    if np.any(target_r < -1e-14) or np.any(target_r > grid.R + 1e-14):
        raise ValueError("Target radii must lie within [0, R].")

    T0 = center_reconstruction(T_C)
    C0 = center_reconstruction(C)
    if abs(t) <= 1e-14:
        Ts, Cs = params.T0_C, params.C0
    else:
        Ts, Cs, _ = coupled_surface_values(
            T_C, C, t, grid, env, params, config
        )
    r_aug = np.concatenate([[0.0], grid.centers, [grid.R]])
    T_aug = np.concatenate([[T0], T_C, [Ts]])
    C_aug = np.concatenate([[C0], C, [Cs]])
    return np.interp(target_r, r_aug, T_aug), np.interp(target_r, r_aug, C_aug)



# -----------------------------------------------------------------------------
# Full-domain maximum and event definition
# -----------------------------------------------------------------------------

def augmented_cmax(
    t: float,
    y: Array,
    grid: FVMGrid,
    env: LongTermEnvironment,
    params: Q3Parameters,
    rhs_config: SolverConfig,
) -> tuple[float, float]:
    """Discrete full-domain Cmax using center reconstruction, all cells, and surface.

    Because output interpolation is piecewise linear between these augmented
    locations, its maximum is attained at one of these values. This avoids
    defining the event from the coarser 0.1 cm result grid.
    """
    N = grid.N
    T = np.asarray(y[:N], dtype=float)
    C = np.asarray(y[N:], dtype=float)

    C0 = center_reconstruction(C)
    _, Cs, _ = coupled_surface_values(T, C, float(t), grid, env, params, rhs_config)

    vals = np.concatenate([[C0], C, [Cs]])
    locs = np.concatenate([[0.0], grid.centers, [grid.R]])
    j = int(np.argmax(vals))
    return float(vals[j]), float(locs[j])


def make_drying_event(
    grid: FVMGrid,
    env: LongTermEnvironment,
    params: Q3Parameters,
    rhs_config: SolverConfig,
):
    def event(t: float, y: Array) -> float:
        cmax, _ = augmented_cmax(t, y, grid, env, params, rhs_config)
        return cmax - params.threshold

    event.terminal = True
    event.direction = -1.0
    return event


# -----------------------------------------------------------------------------
# Long-time segment/output helpers
# -----------------------------------------------------------------------------

def _segment_edges(env: LongTermEnvironment, config: Q3SolverConfig) -> Array:
    end = float(config.max_time_s)
    data_end = env.data_end_s
    if end <= 0:
        raise ValueError("max_time_s must be positive.")
    if config.long_segment_s <= 0:
        raise ValueError("long_segment_s must be positive.")

    edges = [0.0]
    if config.segment_at_data_knots:
        for x in env.t:
            x = float(x)
            if 0.0 < x < min(data_end, end) - 1e-12:
                edges.append(x)
    if data_end < end - 1e-12:
        edges.append(data_end)
        x = data_end
        while x < end - 1e-12:
            x = min(end, x + config.long_segment_s)
            edges.append(x)
    else:
        edges.append(end)

    arr = np.array(sorted(set(round(x, 12) for x in edges)), dtype=float)
    if arr[-1] < end - 1e-10:
        arr = np.append(arr, end)
    return arr


def _regular_times_between(
    a: float,
    b: float,
    dt: float,
    include_left: bool,
) -> Array:
    if dt <= 0:
        raise ValueError("output_dt must be positive.")
    n0 = int(math.ceil((a - 1e-10) / dt))
    n1 = int(math.floor((b + 1e-10) / dt))
    if n1 < n0:
        return np.array([], dtype=float)
    tt = np.arange(n0, n1 + 1, dtype=float) * dt
    tt = tt[(tt >= a - 1e-10) & (tt <= b + 1e-10)]
    if not include_left:
        tt = tt[tt > a + 1e-10]
    return tt


def _refine_event(
    ivp,
    t_event: float,
    grid: FVMGrid,
    env: LongTermEnvironment,
    params: Q3Parameters,
    config: Q3SolverConfig,
    rhs_config: SolverConfig,
) -> tuple[float, Array, dict]:
    """Locally re-integrate before the detected root with tighter controls."""
    if not config.refine_event:
        return float(t_event), np.asarray(ivp.y_events[0][0], dtype=float), {
            "refined": False, "baseline_event_s": float(t_event),
        }
    if ivp.sol is None:
        raise RuntimeError("Event refinement requires dense_output=True.")

    left = max(float(ivp.t[0]), float(t_event) - config.event_refine_window_s)
    y_left = np.asarray(ivp.sol(left), dtype=float)

    # Integrate slightly beyond the baseline event. The tighter run terminates at
    # its own first downward crossing.
    right = float(t_event) + config.event_refine_window_s
    rhs = lambda tt, yy: joint_rhs(tt, yy, grid, env, params, rhs_config, None)
    event = make_drying_event(grid, env, params, rhs_config)
    sparsity = jacobian_sparsity(grid.N) if config.method.upper() in {"BDF", "RADAU"} else None

    ref = solve_ivp(
        rhs,
        (left, right),
        y_left,
        method=config.method,
        rtol=config.event_rtol,
        atol=config.event_atol,
        max_step=config.event_max_step,
        jac_sparsity=sparsity,
        events=event,
    )
    if not ref.success:
        raise RuntimeError(f"Local event refinement failed: {ref.message}")
    if not ref.t_events or ref.t_events[0].size == 0:
        raise RuntimeError("Local event refinement did not rediscover the terminal crossing.")

    tr = float(ref.t_events[0][0])
    yr = np.asarray(ref.y_events[0][0], dtype=float)
    return tr, yr, {
        "refined": True,
        "baseline_event_s": float(t_event),
        "refined_event_s": tr,
        "refinement_shift_s": tr - float(t_event),
        "refine_nfev": int(ref.nfev),
        "refine_njev": int(getattr(ref, "njev", 0) or 0),
        "refine_nlu": int(getattr(ref, "nlu", 0) or 0),
    }


# -----------------------------------------------------------------------------
# Production long-time solver
# -----------------------------------------------------------------------------

def solve_q3(
    env_path: str | Path,
    N: int = 640,
    config: Optional[Q3SolverConfig] = None,
    params: Optional[Q3Parameters] = None,
    boundary_strategy: BoundaryStrategy = "last_hour_mean",
) -> Q3Solution:
    """Q3 production solver: Q2 joint FVM+BDF + long boundary + terminal event."""
    config = config or Q3SolverConfig()
    params = params or Q3Parameters()
    env = LongTermEnvironment.from_excel(env_path, boundary_strategy)
    grid = FVMGrid.build(N, params.R)
    rhs_config = config.rhs_config()

    n_r = int(round(params.R / config.output_dr))
    if not np.isclose(n_r * config.output_dr, params.R, atol=1e-12, rtol=0):
        raise ValueError("R must be an integer multiple of output_dr.")
    r_out = np.linspace(0.0, params.R, n_r + 1)

    y = np.concatenate([
        np.full(N, params.T0_C, dtype=float),
        np.full(N, params.C0, dtype=float),
    ])

    edges = _segment_edges(env, config)
    sparsity = jacobian_sparsity(N) if config.method.upper() in {"BDF", "RADAU"} else None
    event = make_drying_event(grid, env, params, rhs_config)

    surface_counter: dict = {}
    totals = {"segments": 0, "nfev": 0, "njev": 0, "nlu": 0}

    t_chunks: list[Array] = []
    T_chunks: list[Array] = []
    C_chunks: list[Array] = []

    state_t = [0.0]
    T_state = [y[:N].copy()]
    C_state = [y[N:].copy()]

    t_dry = None
    y_dry = None
    event_meta: dict = {}

    for s in range(len(edges) - 1):
        a, b = float(edges[s]), float(edges[s + 1])
        tev = _regular_times_between(a, b, config.output_dt, include_left=(s == 0))
        max_step = config.max_step_data if a < env.data_end_s - 1e-12 else config.max_step_long

        rhs = lambda tt, yy: joint_rhs(
            tt, yy, grid, env, params, rhs_config, surface_counter
        )
        ivp = solve_ivp(
            rhs,
            (a, b),
            y,
            method=config.method,
            t_eval=tev if tev.size else None,
            rtol=config.rtol,
            atol=config.atol,
            max_step=max_step,
            jac_sparsity=sparsity,
            events=event,
            dense_output=config.refine_event,
        )
        if not ivp.success:
            raise RuntimeError(f"Q3 BDF failed on [{a},{b}] s: {ivp.message}")

        totals["segments"] += 1
        totals["nfev"] += int(ivp.nfev)
        totals["njev"] += int(getattr(ivp, "njev", 0) or 0)
        totals["nlu"] += int(getattr(ivp, "nlu", 0) or 0)

        # t_eval output returned only up to a terminal event.
        if ivp.t.size:
            segT = []
            segC = []
            for jj, tt in enumerate(ivp.t):
                Tj = ivp.y[:N, jj]
                Cj = ivp.y[N:, jj]
                Tsamp, Csamp = sample_state_pair(
                    Tj, Cj, float(tt), r_out, grid, env, params, rhs_config
                )
                segT.append(Tsamp)
                segC.append(Csamp)
            t_chunks.append(np.asarray(ivp.t))
            T_chunks.append(np.asarray(segT))
            C_chunks.append(np.asarray(segC))

        if ivp.t_events and ivp.t_events[0].size:
            t0_event = float(ivp.t_events[0][0])
            t_dry, y_dry, event_meta = _refine_event(
                ivp, t0_event, grid, env, params, config, rhs_config
            )
            state_t.append(float(t_dry))
            T_state.append(y_dry[:N].copy())
            C_state.append(y_dry[N:].copy())
            break

        # No event in this segment: continue from the true segment endpoint.
        # If t_eval did not include b, use dense output or re-evaluate with a tiny
        # endpoint request. In the standard Q3 grids b is always a 60 s multiple.
        if ivp.t.size == 0 or abs(float(ivp.t[-1]) - b) > 1e-8:
            # Run a minimal endpoint solve from the last returned state/time.
            ta = float(ivp.t[-1]) if ivp.t.size else a
            ya = ivp.y[:, -1].copy() if ivp.t.size else y.copy()
            tail = solve_ivp(
                rhs, (ta, b), ya, method=config.method,
                t_eval=[b], rtol=config.rtol, atol=config.atol, max_step=max_step,
                jac_sparsity=sparsity,
            )
            if not tail.success:
                raise RuntimeError(f"Endpoint continuation failed on [{ta},{b}] s.")
            y = tail.y[:, -1].copy()
            totals["nfev"] += int(tail.nfev)
            totals["njev"] += int(getattr(tail, "njev", 0) or 0)
            totals["nlu"] += int(getattr(tail, "nlu", 0) or 0)
        else:
            y = ivp.y[:, -1].copy()

        state_t.append(b)
        T_state.append(y[:N].copy())
        C_state.append(y[N:].copy())

    if t_dry is None or y_dry is None:
        cmax_last, rmax_last = augmented_cmax(
            state_t[-1],
            np.concatenate([T_state[-1], C_state[-1]]),
            grid, env, params, rhs_config,
        )
        raise RuntimeError(
            f"Drying threshold was not reached by max_time_s={config.max_time_s:.6g}. "
            f"Final monitored Cmax={cmax_last:.6g} at r={100*rmax_last:.4f} cm. "
            "Increase max_time_s only if the model assumptions justify it."
        )

    if not t_chunks:
        raise RuntimeError("No regular output samples were produced.")

    tout = np.concatenate(t_chunks)
    T_out = np.vstack(T_chunks)
    C_out = np.vstack(C_chunks)

    # Remove accidental duplicates from segment boundaries.
    order = np.argsort(tout)
    tout, T_out, C_out = tout[order], T_out[order], C_out[order]
    _, unique_idx = np.unique(np.round(tout, 10), return_index=True)
    unique_idx = np.sort(unique_idx)
    tout, T_out, C_out = tout[unique_idx], T_out[unique_idx], C_out[unique_idx]

    # result3 stops at the last regular 60 s sample not exceeding tdry.
    t_last = config.output_dt * math.floor((t_dry + 1e-10) / config.output_dt)
    keep = tout <= t_last + 1e-8
    tout, T_out, C_out = tout[keep], T_out[keep], C_out[keep]

    expected = np.arange(0.0, t_last + 0.5 * config.output_dt, config.output_dt)
    if len(tout) != len(expected) or not np.allclose(tout, expected, atol=1e-8, rtol=0):
        raise RuntimeError("Regular Q3 output times are incomplete or misaligned.")

    Cmax_dry, rmax_dry = augmented_cmax(
        t_dry, y_dry, grid, env, params, rhs_config
    )

    totals.update({
        "surface_calls": int(surface_counter.get("calls", 0)),
        "surface_total_iterations": int(surface_counter.get("iterations", 0)),
        "surface_max_iterations": int(surface_counter.get("max_iterations", 0)),
        "surface_mean_iterations": (
            float(surface_counter.get("iterations", 0))
            / max(1, int(surface_counter.get("calls", 0)))
        ),
        "boundary_strategy": boundary_strategy,
        "surface_picard_success_calls": int(surface_counter.get("picard_success", 0)),
        "surface_brent_fallback_calls": int(surface_counter.get("brent_fallback", 0)),
        "T_long_C": env.stats.T_long_C,
        "C_long": env.stats.C_long,
        **event_meta,
    })

    return Q3Solution(
        grid=grid, env=env, params=params, config=config,
        t=tout, r=r_out, T=T_out, C=C_out,
        state_t=np.asarray(state_t, dtype=float),
        T_state=np.asarray(T_state, dtype=float),
        C_state=np.asarray(C_state, dtype=float),
        t_dry_s=float(t_dry),
        T_dry_state=np.asarray(y_dry[:N], dtype=float),
        C_dry_state=np.asarray(y_dry[N:], dtype=float),
        Cmax_dry=float(Cmax_dry),
        rmax_dry_m=float(rmax_dry),
        solver_stats=totals,
    )


# -----------------------------------------------------------------------------
# Lightweight event-only solve for convergence/sensitivity studies
# -----------------------------------------------------------------------------

def find_drying_time(
    env_path: str | Path,
    N: int = 320,
    config: Optional[Q3SolverConfig] = None,
    params: Optional[Q3Parameters] = None,
    boundary_strategy: BoundaryStrategy = "last_hour_mean",
) -> DryingTimeResult:
    config = config or Q3SolverConfig()
    params = params or Q3Parameters()
    env = LongTermEnvironment.from_excel(env_path, boundary_strategy)
    grid = FVMGrid.build(N, params.R)
    rhs_config = config.rhs_config()
    sparsity = jacobian_sparsity(N) if config.method.upper() in {"BDF", "RADAU"} else None
    event = make_drying_event(grid, env, params, rhs_config)

    y = np.concatenate([
        np.full(N, params.T0_C, dtype=float),
        np.full(N, params.C0, dtype=float),
    ])
    edges = _segment_edges(env, config)
    totals = {"segments": 0, "nfev": 0, "njev": 0, "nlu": 0}
    surface_counter: dict = {}

    for s in range(len(edges) - 1):
        a, b = float(edges[s]), float(edges[s + 1])
        max_step = config.max_step_data if a < env.data_end_s - 1e-12 else config.max_step_long
        rhs = lambda tt, yy: joint_rhs(
            tt, yy, grid, env, params, rhs_config, surface_counter
        )
        ivp = solve_ivp(
            rhs, (a, b), y,
            method=config.method,
            rtol=config.rtol, atol=config.atol,
            max_step=max_step,
            jac_sparsity=sparsity,
            events=event,
            dense_output=config.refine_event,
        )
        if not ivp.success:
            raise RuntimeError(f"Q3 event search failed on [{a},{b}] s: {ivp.message}")

        totals["segments"] += 1
        totals["nfev"] += int(ivp.nfev)
        totals["njev"] += int(getattr(ivp, "njev", 0) or 0)
        totals["nlu"] += int(getattr(ivp, "nlu", 0) or 0)

        if ivp.t_events and ivp.t_events[0].size:
            te0 = float(ivp.t_events[0][0])
            te, ye, meta = _refine_event(
                ivp, te0, grid, env, params, config, rhs_config
            )
            cmax, rmax = augmented_cmax(te, ye, grid, env, params, rhs_config)
            totals.update({
                "surface_calls": int(surface_counter.get("calls", 0)),
                "surface_total_iterations": int(surface_counter.get("iterations", 0)),
                "surface_max_iterations": int(surface_counter.get("max_iterations", 0)),
                "surface_mean_iterations": (
                    float(surface_counter.get("iterations", 0))
                    / max(1, int(surface_counter.get("calls", 0)))
                ),
                "boundary_strategy": boundary_strategy,
                "surface_picard_success_calls": int(surface_counter.get("picard_success", 0)),
                "surface_brent_fallback_calls": int(surface_counter.get("brent_fallback", 0)),
                **meta,
            })
            return DryingTimeResult(
                N=N,
                t_dry_s=float(te),
                t_dry_h=float(te / 3600.0),
                Cmax=float(cmax),
                rmax_m=float(rmax),
                solver_stats=totals,
            )

        y = ivp.y[:, -1].copy()

    cmax, rmax = augmented_cmax(
        edges[-1], y, grid, env, params, rhs_config
    )
    raise RuntimeError(
        f"No drying event by {config.max_time_s/3600:.3f} h; "
        f"Cmax={cmax:.6g} at r={100*rmax:.4f} cm."
    )


# -----------------------------------------------------------------------------
# Diagnostics / tables
# -----------------------------------------------------------------------------

def environment_summary(env: LongTermEnvironment) -> pd.DataFrame:
    s = env.stats
    return pd.DataFrame([{
        "strategy": s.strategy,
        "data_end_h": s.data_end_s / 3600.0,
        "window_start_h": s.window_start_s / 3600.0,
        "window_end_h": s.window_end_s / 3600.0,
        "n_points": s.n_points,
        "T_long_C": s.T_long_C,
        "T_std_C": s.T_std_C,
        "T_slope_C_per_h": s.T_slope_C_per_h,
        "C_long": s.C_long,
        "C_std": s.C_std,
        "C_slope_per_h": s.C_slope_per_h,
    }])


def drying_event_summary(sol: Q3Solution) -> dict:
    rhs_cfg = sol.config.rhs_config()
    y = np.concatenate([sol.T_dry_state, sol.C_dry_state])
    cmax, rmax = augmented_cmax(
        sol.t_dry_s, y, sol.grid, sol.env, sol.params, rhs_cfg
    )
    T_event, C_event = sample_state_pair(
        sol.T_dry_state, sol.C_dry_state, sol.t_dry_s,
        np.array([0.0, 0.005, 0.010, 0.015, 0.020]),
        sol.grid, sol.env, sol.params, rhs_cfg,
    )
    return {
        "t_dry_s": sol.t_dry_s,
        "t_dry_h": sol.t_dry_s / 3600.0,
        "Cmax_at_event": cmax,
        "threshold": sol.params.threshold,
        "g_event": cmax - sol.params.threshold,
        "rmax_cm": 100.0 * rmax,
        "C_center": float(C_event[0]),
        "C_surface": float(C_event[-1]),
        "T_center_C": float(T_event[0]),
        "T_surface_C": float(T_event[-1]),
    }


def cmax_location_history(sol: Q3Solution) -> pd.DataFrame:
    rhs_cfg = sol.config.rhs_config()
    rows = []
    for t, T, C in zip(sol.state_t, sol.T_state, sol.C_state):
        y = np.concatenate([T, C])
        cmax, rmax = augmented_cmax(
            float(t), y, sol.grid, sol.env, sol.params, rhs_cfg
        )
        rows.append({
            "time_s": float(t),
            "time_h": float(t) / 3600.0,
            "Cmax": cmax,
            "g": cmax - sol.params.threshold,
            "rmax_m": rmax,
            "rmax_cm": 100.0 * rmax,
        })
    return pd.DataFrame(rows)


def physical_diagnostics_q3(sol: Q3Solution) -> dict:
    hist = cmax_location_history(sol)
    return {
        "T_min_C": float(np.min(sol.T)),
        "T_max_C": float(np.max(sol.T)),
        "C_min": float(np.min(sol.C)),
        "C_max": float(np.max(sol.C)),
        "all_C_positive": bool(np.all(sol.C > 0.0)),
        "regular_output_C_nonincreasing_time_fraction_by_radius": float(
            np.mean(np.all(np.diff(sol.C, axis=0) <= 1e-8, axis=0))
        ),
        "state_Cmax_nonincreasing": bool(
            np.all(np.diff(hist["Cmax"].to_numpy()) <= 1e-8)
        ),
        "rmax_at_center_fraction": float(
            np.mean(np.isclose(hist["rmax_m"].to_numpy(), 0.0, atol=1e-12, rtol=0))
        ),
        "event_rmax_cm": 100.0 * sol.rmax_dry_m,
        "surface_picard_mean_iterations": sol.solver_stats.get("surface_mean_iterations"),
        "surface_picard_max_iterations": sol.solver_stats.get("surface_max_iterations"),
    }


def paper_table5(sol: Q3Solution) -> pd.DataFrame:
    radii = np.array([0.0, 0.005, 0.010, 0.015, 0.020])
    cols = ["0", "0.5", "1.0", "1.5", "2.0"]

    scheduled = np.arange(6.0 * 3600.0, sol.t_dry_s - 1e-8, 6.0 * 3600.0)
    rows = []
    labels = []

    # Scheduled times are exact multiples of 60 s, hence present in regular output.
    for tt in scheduled:
        j = np.where(np.isclose(sol.t, tt, atol=1e-8, rtol=0))[0]
        if j.size == 0:
            raise RuntimeError(f"Scheduled table time {tt} s is missing from Q3 output.")
        ir = [int(np.where(np.isclose(sol.r, r, atol=1e-12, rtol=0))[0][0]) for r in radii]
        rows.append(sol.C[int(j[0]), ir])
        labels.append(f"{tt/3600:.0f}")

    _, Cevt = sample_state_pair(
        sol.T_dry_state, sol.C_dry_state, sol.t_dry_s,
        radii, sol.grid, sol.env, sol.params, sol.config.rhs_config()
    )
    rows.append(Cevt)
    labels.append(f"Dry end ({sol.t_dry_s/3600:.6f} h)")

    df = pd.DataFrame(np.asarray(rows), index=labels, columns=cols)
    df.index.name = "Time (h)"
    df.columns.name = "Radius (cm)"
    return df


def event_grid_convergence(
    env_path: str | Path,
    N_values: Sequence[int] = (320, 640, 1280),
    config: Optional[Q3SolverConfig] = None,
    boundary_strategy: BoundaryStrategy = "last_hour_mean",
) -> tuple[pd.DataFrame, dict[int, DryingTimeResult]]:
    config = config or Q3SolverConfig()
    results: dict[int, DryingTimeResult] = {}
    for N in N_values:
        results[int(N)] = find_drying_time(
            env_path, N=int(N), config=config, boundary_strategy=boundary_strategy
        )
    rows = []
    for Na, Nb in zip(N_values[:-1], N_values[1:]):
        a, b = results[int(Na)], results[int(Nb)]
        rows.append({
            "coarse_N": int(Na),
            "fine_N": int(Nb),
            "tdry_coarse_h": a.t_dry_h,
            "tdry_fine_h": b.t_dry_h,
            "abs_difference_s": abs(b.t_dry_s - a.t_dry_s),
            "abs_difference_min": abs(b.t_dry_s - a.t_dry_s) / 60.0,
        })
    return pd.DataFrame(rows), results


def event_time_sensitivity(
    env_path: str | Path,
    N: int = 320,
    baseline: Optional[Q3SolverConfig] = None,
    boundary_strategy: BoundaryStrategy = "last_hour_mean",
) -> pd.DataFrame:
    baseline = baseline or Q3SolverConfig()
    tighter = replace(
        baseline,
        rtol=min(baseline.rtol, 2e-10),
        atol=min(baseline.atol, 2e-12),
        max_step_long=min(baseline.max_step_long, 15.0),
        event_max_step=min(baseline.event_max_step, 2.0),
    )
    a = find_drying_time(env_path, N=N, config=baseline, boundary_strategy=boundary_strategy)
    b = find_drying_time(env_path, N=N, config=tighter, boundary_strategy=boundary_strategy)
    return pd.DataFrame([
        {"case": "baseline", "tdry_h": a.t_dry_h, "tdry_s": a.t_dry_s},
        {"case": "tighter", "tdry_h": b.t_dry_h, "tdry_s": b.t_dry_s},
        {"case": "difference", "tdry_h": abs(b.t_dry_s-a.t_dry_s)/3600.0,
         "tdry_s": abs(b.t_dry_s-a.t_dry_s)},
    ])


def boundary_strategy_sensitivity(
    env_path: str | Path,
    N: int = 320,
    config: Optional[Q3SolverConfig] = None,
    strategies: Sequence[BoundaryStrategy] = (
        "last_hour_mean", "last_value", "last_1p5h_mean"
    ),
) -> pd.DataFrame:
    config = config or Q3SolverConfig()
    rows = []
    for strategy in strategies:
        env = LongTermEnvironment.from_excel(env_path, strategy)
        res = find_drying_time(env_path, N=N, config=config, boundary_strategy=strategy)
        rows.append({
            "strategy": strategy,
            "T_long_C": env.stats.T_long_C,
            "C_long": env.stats.C_long,
            "tdry_h": res.t_dry_h,
            "tdry_s": res.t_dry_s,
        })
    df = pd.DataFrame(rows)
    main = float(df.loc[df["strategy"] == "last_hour_mean", "tdry_s"].iloc[0])
    spread = float(df["tdry_s"].max() - df["tdry_s"].min())
    df["relative_spread_vs_main"] = spread / main
    return df


def event_local_trace(
    sol: Q3Solution,
    before_s: float = 600.0,
    after_s: float = 600.0,
    dt: float = 5.0,
) -> pd.DataFrame:
    """Reintegrate around tdry without a terminal event to display a stable crossing."""
    if dt <= 0 or before_s < 0 or after_s < 0:
        raise ValueError("Invalid local trace settings.")

    target_start = sol.t_dry_s - before_s
    eligible = np.where(sol.state_t <= target_start + 1e-10)[0]
    idx = int(eligible[-1]) if eligible.size else 0
    t0 = float(sol.state_t[idx])
    y0 = np.concatenate([sol.T_state[idx], sol.C_state[idx]])

    t1 = sol.t_dry_s + after_s
    times = np.arange(
        math.ceil(t0 / dt) * dt,
        math.floor(t1 / dt) * dt + 0.5 * dt,
        dt,
    )
    times = times[(times >= t0 - 1e-10) & (times <= t1 + 1e-10)]

    rhs_cfg = sol.config.rhs_config()
    rhs = lambda tt, yy: joint_rhs(
        tt, yy, sol.grid, sol.env, sol.params, rhs_cfg, None
    )
    sparsity = jacobian_sparsity(sol.grid.N) if sol.config.method.upper() in {"BDF", "RADAU"} else None
    ivp = solve_ivp(
        rhs, (t0, t1), y0,
        method=sol.config.method,
        t_eval=times,
        rtol=sol.config.event_rtol,
        atol=sol.config.event_atol,
        max_step=min(sol.config.event_max_step, dt),
        jac_sparsity=sparsity,
    )
    if not ivp.success:
        raise RuntimeError(f"Local event trace failed: {ivp.message}")

    rows = []
    for j, tt in enumerate(ivp.t):
        cmax, rmax = augmented_cmax(
            float(tt), ivp.y[:, j], sol.grid, sol.env, sol.params, rhs_cfg
        )
        rows.append({
            "time_s": float(tt),
            "time_h": float(tt)/3600.0,
            "relative_to_event_s": float(tt-sol.t_dry_s),
            "Cmax": cmax,
            "g": cmax-sol.params.threshold,
            "rmax_cm": 100.0*rmax,
        })
    # Add the refined root explicitly.
    rows.append({
        "time_s": sol.t_dry_s,
        "time_h": sol.t_dry_s/3600.0,
        "relative_to_event_s": 0.0,
        "Cmax": sol.Cmax_dry,
        "g": sol.Cmax_dry-sol.params.threshold,
        "rmax_cm": 100.0*sol.rmax_dry_m,
    })
    return pd.DataFrame(rows).sort_values("time_s").drop_duplicates("time_s")


def long_time_balance_diagnostics(sol: Q3Solution) -> dict:
    """Semi-discrete instantaneous balances + moisture integral balance.

    The heat check is instantaneous because rho(C)*cp(C) varies with state;
    we intentionally do NOT reuse Q1's constant-capacity cumulative-energy test.
    """
    rhs_cfg = sol.config.rhs_config()

    heat_rel = []
    moist_rel = []
    M = []
    for t, T, C in zip(sol.state_t, sol.T_state, sol.C_state):
        y = np.concatenate([T, C])
        dy = joint_rhs(float(t), y, sol.grid, sol.env, sol.params, rhs_cfg, None)
        dT = dy[:sol.grid.N]
        dC = dy[sol.grid.N:]
        Ts, Cs, _ = coupled_surface_values(
            T, C, float(t), sol.grid, sol.env, sol.params, rhs_cfg
        )
        qR = sol.params.h * (Ts - float(sol.env.T_inf(t)))
        JR = sol.params.hm * (Cs - float(sol.env.C_inf(t)))
        cap = density(C) * heat_capacity(C)

        heat_lhs = float(np.sum(cap * sol.grid.volumes * dT) + sol.grid.areas[-1] * qR)
        heat_scale = max(
            abs(float(np.sum(cap * sol.grid.volumes * dT))),
            abs(float(sol.grid.areas[-1] * qR)),
            1e-30,
        )
        moist_lhs = float(np.sum(sol.grid.volumes * dC) + sol.grid.areas[-1] * JR)
        moist_scale = max(
            abs(float(np.sum(sol.grid.volumes * dC))),
            abs(float(sol.grid.areas[-1] * JR)),
            1e-30,
        )
        heat_rel.append(abs(heat_lhs) / heat_scale)
        moist_rel.append(abs(moist_lhs) / moist_scale)
        M.append(float(np.sum(sol.grid.volumes * C)))

    # Use regular 60 s output for accurate boundary-flux quadrature.
    Jout = sol.params.hm * (sol.C[:, -1] - np.asarray(sol.env.C_inf(sol.t), dtype=float))
    cumulative_out = cumulative_trapezoid(
        sol.grid.areas[-1] * Jout, sol.t, initial=0.0
    )
    # Evaluate cumulative flux at the state-monitoring times that lie on/before
    # the final regular result time. The exact event itself may be non-60 s.
    st = np.asarray(sol.state_t)
    mask = st <= sol.t[-1] + 1e-10
    st2 = st[mask]
    M2 = np.asarray(M)[mask]
    out2 = np.interp(st2, sol.t, cumulative_out)
    residual = (M2 - M2[0]) + out2
    scale = max(abs(M2[0]-M2[-1]), abs(out2[-1]), 1e-30)

    return {
        "heat_instantaneous_max_relative": float(np.max(heat_rel)),
        "moisture_instantaneous_max_relative": float(np.max(moist_rel)),
        "moisture_integral_global_normalized_max": float(np.max(np.abs(residual))/scale),
        "moisture_integral_final_relative": float(abs(residual[-1])/scale),
        "state_time_s": st2,
        "moisture_integral_residual": residual,
    }


# -----------------------------------------------------------------------------
# Independent fixed-step backward-Euler + block-Picard event reference
# -----------------------------------------------------------------------------

def find_drying_time_block_picard(
    env_path: str | Path,
    N: int = 80,
    dt: float = 600.0,
    params: Optional[Q3Parameters] = None,
    boundary_strategy: BoundaryStrategy = "last_hour_mean",
    max_time_s: float = 15.0 * 24.0 * 3600.0,
    tol_T: float = 1e-8,
    tol_C: float = 1e-10,
    max_iter: int = 50,
    relaxation: float = 1.0,
    face_average: FaceAverage = "harmonic",
) -> BlockPicardEventResult:
    """Coarse independent BE+block-Picard reference for tdry convergence analysis."""
    params = params or Q3Parameters()
    env = LongTermEnvironment.from_excel(env_path, boundary_strategy)
    grid = FVMGrid.build(N, params.R)
    surf_cfg = SolverConfig(face_average=face_average)

    T = np.full(N, params.T0_C, dtype=float)
    C = np.full(N, params.C0, dtype=float)
    y = np.concatenate([T, C])
    g_old = augmented_cmax(0.0, y, grid, env, params, surf_cfg)[0] - params.threshold

    nsteps = int(math.ceil(max_time_s / dt))
    iterations = []

    for n in range(nsteps):
        tnew = min((n + 1) * dt, max_time_s)
        dt_n = tnew - n * dt
        if dt_n <= 0:
            break

        Told, Cold = T.copy(), C.copy()
        Tg, Cg = Told.copy(), Cold.copy()
        converged = False

        for m in range(1, max_iter + 1):
            check_material_state(Tg, Cg)

            Ts_g, Cs_g, _ = coupled_surface_values(
                Tg, Cg, tnew, grid, env, params, surf_cfg
            )
            cap = density(Cg) * heat_capacity(Cg)
            kf = _k_faces(Cg, face_average)
            ks = float(thermal_conductivity(Cs_g))
            Tlin = _be_diffusion_solve(
                Told, cap, kf, ks, params.h,
                float(env.T_inf(tnew)), grid, dt_n,
            )
            Tnew = relaxation*Tlin + (1.0-relaxation)*Tg

            Ts_aux, Cs_aux, _ = coupled_surface_values(
                Tnew, Cg, tnew, grid, env, params, surf_cfg
            )
            Df = _D_faces(Cg, Tnew, face_average)
            Ds = float(moisture_diffusivity(Cs_aux, Ts_aux))
            Clin = _be_diffusion_solve(
                Cold, np.ones(N), Df, Ds, params.hm,
                float(env.C_inf(tnew)), grid, dt_n,
            )
            Cnew = relaxation*Clin + (1.0-relaxation)*Cg

            if np.any(Cnew <= 0) or not np.all(np.isfinite(Tnew)) or not np.all(np.isfinite(Cnew)):
                raise FloatingPointError(f"Invalid block-Picard state at t={tnew} s.")

            eT = float(np.max(np.abs(Tnew-Tg)))
            eC = float(np.max(np.abs(Cnew-Cg)))
            Tg, Cg = Tnew, Cnew
            if eT <= tol_T and eC <= tol_C:
                converged = True
                iterations.append(m)
                break

        if not converged:
            raise RuntimeError(
                f"Block Picard failed at t={tnew:.3f} s after {max_iter} iterations."
            )

        T, C = Tg, Cg
        ynew = np.concatenate([T, C])
        g_new = augmented_cmax(tnew, ynew, grid, env, params, surf_cfg)[0] - params.threshold

        if g_old > 0.0 and g_new <= 0.0:
            # Auxiliary first-order reference only: linear root interpolation
            # inside the final fixed BE interval.
            frac = g_old / (g_old - g_new)
            te = (tnew - dt_n) + frac * dt_n
            return BlockPicardEventResult(
                N=N,
                dt_s=dt,
                t_dry_s=float(te),
                t_dry_h=float(te/3600.0),
                max_picard_iterations=int(max(iterations)),
                mean_picard_iterations=float(np.mean(iterations)),
            )

        g_old = g_new

    raise RuntimeError("Block-Picard reference did not reach the drying threshold.")


# -----------------------------------------------------------------------------
# Result writer
# -----------------------------------------------------------------------------

def write_result3(
    sol: Q3Solution,
    output_path: str | Path,
    template_path: Optional[str | Path] = None,
    decimals: int = 4,
) -> Path:
    """Write only the required 60 s x 0.1 cm moisture table.

    The exact non-60-s tdry is intentionally NOT inserted as an extra time row.
    Report it separately in the paper/table 5.
    """
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    if template_path is not None:
        wb = load_workbook(template_path)
        if "水分浓度" in wb.sheetnames:
            ws = wb["水分浓度"]
        else:
            ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "水分浓度"

    ws.cell(row=1, column=1, value="时间/s")
    for j, r in enumerate(sol.r, start=2):
        ws.cell(row=1, column=j, value=round(float(r*100.0), 10))

    for i, tt in enumerate(sol.t, start=2):
        ws.cell(row=i, column=1, value=int(round(float(tt))))
        for j in range(sol.C.shape[1]):
            ws.cell(row=i, column=j+2, value=round(float(sol.C[i-2, j]), decimals))

    wb.save(output_path)
    return output_path


# -----------------------------------------------------------------------------
# CLI
# -----------------------------------------------------------------------------

def _build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        description="Q3 drying solver: Q2 coupled FVM+BDF + long boundary + terminal event"
    )
    p.add_argument("--env", required=True, help="Path to Attachment 1 xlsx")
    p.add_argument("--output", required=True, help="Path to result3.xlsx")
    p.add_argument("--template", default=None, help="Optional official result3.xlsx template")
    p.add_argument("--N", type=int, default=640)
    p.add_argument(
        "--boundary-strategy",
        choices=["last_hour_mean", "last_value", "last_1p5h_mean"],
        default="last_hour_mean",
    )
    p.add_argument("--rtol", type=float, default=1e-9)
    p.add_argument("--atol", type=float, default=1e-11)
    p.add_argument("--max-step-data", type=float, default=5.0)
    p.add_argument("--max-step-long", type=float, default=60.0)
    p.add_argument("--event-max-step", type=float, default=5.0)
    p.add_argument("--max-time-hours", type=float, default=15.0*24.0)
    p.add_argument("--no-event-refine", action="store_true")
    return p


def main() -> None:
    args = _build_parser().parse_args()
    cfg = Q3SolverConfig(
        rtol=args.rtol,
        atol=args.atol,
        max_step_data=args.max_step_data,
        max_step_long=args.max_step_long,
        event_max_step=args.event_max_step,
        max_time_s=args.max_time_hours*3600.0,
        refine_event=not args.no_event_refine,
    )
    sol = solve_q3(
        args.env, N=args.N, config=cfg,
        boundary_strategy=args.boundary_strategy,
    )
    print("Environment extension:")
    print(environment_summary(sol.env).to_string(index=False))
    print("Drying event:")
    print(drying_event_summary(sol))
    print("Physical diagnostics:")
    print(physical_diagnostics_q3(sol))
    out = write_result3(sol, args.output, template_path=args.template)
    print("Saved:", out)


if __name__ == "__main__":
    main()
