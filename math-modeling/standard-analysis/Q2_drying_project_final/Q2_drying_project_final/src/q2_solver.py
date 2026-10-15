from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Iterable, Literal, Optional, Sequence
import argparse
import math

import numpy as np
import pandas as pd
from scipy.integrate import solve_ivp, cumulative_trapezoid
from scipy.sparse import csr_matrix, diags
from scipy.sparse.linalg import spsolve
from openpyxl import load_workbook, Workbook


Array = np.ndarray
FaceAverage = Literal["harmonic", "arithmetic", "midpoint_property"]


@dataclass(frozen=True)
class Q2Parameters:
    R: float = 0.02               # m
    h: float = 25.0               # W/(m^2 K), inherited from Q1
    hm: float = 8.0e-7            # m/s, inherited from Q1
    T0_C: float = 28.0            # degC
    C0: float = 2.55              # kg/kg dry basis


@dataclass(frozen=True)
class SolverConfig:
    method: str = "BDF"
    rtol: float = 1.0e-9
    atol: float = 1.0e-11
    max_step: float = 5.0
    output_dt: float = 1.0
    face_average: FaceAverage = "harmonic"
    segment_at_data_knots: bool = True
    surface_tol_T: float = 1.0e-11
    surface_tol_C: float = 1.0e-12
    surface_max_iter: int = 50
    surface_relaxation: float = 1.0


@dataclass(frozen=True)
class PicardConfig:
    dt: float = 30.0
    tol_T: float = 1.0e-8
    tol_C: float = 1.0e-10
    max_iter: int = 50
    relaxation: float = 1.0
    face_average: FaceAverage = "harmonic"
    one_pass: bool = False


@dataclass(frozen=True)
class FVMGrid:
    N: int
    R: float
    dr: float
    faces: Array
    centers: Array
    areas: Array
    volumes: Array

    @classmethod
    def build(cls, N: int, R: float) -> "FVMGrid":
        if N < 3:
            raise ValueError("N must be at least 3 for second-order boundary reconstruction.")
        if R <= 0:
            raise ValueError("R must be positive.")
        dr = R / N
        faces = np.linspace(0.0, R, N + 1)
        centers = (np.arange(N, dtype=float) + 0.5) * dr
        # Unit axial length. The common length factor cancels in the PDE, while
        # actual cylindrical area/volume makes balance checks transparent.
        areas = 2.0 * np.pi * faces
        volumes = np.pi * (faces[1:] ** 2 - faces[:-1] ** 2)
        return cls(N=N, R=R, dr=dr, faces=faces, centers=centers,
                   areas=areas, volumes=volumes)


@dataclass(frozen=True)
class EnvironmentData:
    t: Array
    temperature: Array
    moisture: Array

    @classmethod
    def from_excel(cls, path: str | Path) -> "EnvironmentData":
        path = Path(path)
        df = pd.read_excel(path)
        required = ["时间", "温度", "水分浓度"]
        missing = [c for c in required if c not in df.columns]
        if missing:
            raise ValueError(f"Missing columns in environment file: {missing}")
        df = df[required].copy()
        if df.isna().any().any():
            raise ValueError("Environment data contain missing values.")
        df = df.sort_values("时间")
        t = df["时间"].to_numpy(dtype=float)
        temp = df["温度"].to_numpy(dtype=float)
        moist = df["水分浓度"].to_numpy(dtype=float)
        if np.any(np.diff(t) <= 0):
            raise ValueError("Environment time values must be strictly increasing.")
        return cls(t=t, temperature=temp, moisture=moist)

    def T_inf(self, t: float | Array) -> Array:
        return np.interp(t, self.t, self.temperature)

    def C_inf(self, t: float | Array) -> Array:
        return np.interp(t, self.t, self.moisture)


@dataclass
class Q2Solution:
    grid: FVMGrid
    env: EnvironmentData
    params: Q2Parameters
    config: SolverConfig
    t: Array                     # required output times, usually every 1 s
    r: Array                     # required output radii, usually every 0.1 cm
    T: Array                     # sampled physical field, shape (len(t), len(r))
    C: Array
    state_t: Array               # full cell-centered snapshots, normally every 60 s
    T_state: Array               # shape (len(state_t), N)
    C_state: Array
    solver_stats: dict = field(default_factory=dict)


@dataclass
class BlockPicardSolution:
    grid: FVMGrid
    env: EnvironmentData
    params: Q2Parameters
    config: PicardConfig
    t: Array
    T_state: Array
    C_state: Array
    iterations: Array


# -----------------------------------------------------------------------------
# Appendix 3 material properties
# -----------------------------------------------------------------------------

def density(C: Array | float) -> Array:
    C = np.asarray(C, dtype=float)
    return 650.0 + 128.0 * C


def heat_capacity(C: Array | float) -> Array:
    C = np.asarray(C, dtype=float)
    return 1450.0 + 2736.0 * C / (C + 1.0)


def thermal_conductivity(C: Array | float) -> Array:
    C = np.asarray(C, dtype=float)
    return 0.21 + 0.38 * C / (C + 1.0)


def moisture_diffusivity(C: Array | float, T_C: Array | float) -> Array:
    C = np.asarray(C, dtype=float)
    T_C = np.asarray(T_C, dtype=float)
    if np.any(C <= 0.0):
        raise FloatingPointError("Moisture content C must remain positive.")
    T_K = T_C + 273.15
    if np.any(T_K <= 0.0):
        raise FloatingPointError("Absolute temperature must remain positive.")
    return 2.4e-3 * np.exp(-0.45 / C) * np.exp(-3850.0 / T_K)


def check_material_state(T_C: Array, C: Array) -> None:
    if not np.all(np.isfinite(T_C)) or not np.all(np.isfinite(C)):
        raise FloatingPointError("Non-finite state encountered.")
    if np.any(C <= 0.0):
        raise FloatingPointError(
            f"Non-positive moisture content encountered: min(C)={np.min(C):.6e}"
        )
    rho = density(C)
    cp = heat_capacity(C)
    k = thermal_conductivity(C)
    D = moisture_diffusivity(C, T_C)
    if np.any(rho <= 0) or np.any(cp <= 0) or np.any(k <= 0) or np.any(D <= 0):
        raise FloatingPointError("A material property became non-positive.")


# -----------------------------------------------------------------------------
# Finite-volume helpers
# -----------------------------------------------------------------------------

def harmonic_mean(a: Array, b: Array) -> Array:
    denom = a + b
    if np.any(denom <= 0):
        raise FloatingPointError("Harmonic-mean denominator must be positive.")
    return 2.0 * a * b / denom


def _k_faces(C: Array, mode: FaceAverage) -> Array:
    kcell = thermal_conductivity(C)
    if mode == "harmonic":
        return harmonic_mean(kcell[:-1], kcell[1:])
    if mode == "arithmetic":
        return 0.5 * (kcell[:-1] + kcell[1:])
    if mode == "midpoint_property":
        return thermal_conductivity(0.5 * (C[:-1] + C[1:]))
    raise ValueError(f"Unknown face-average mode: {mode}")


def _D_faces(C: Array, T_C: Array, mode: FaceAverage) -> Array:
    Dcell = moisture_diffusivity(C, T_C)
    if mode == "harmonic":
        return harmonic_mean(Dcell[:-1], Dcell[1:])
    if mode == "arithmetic":
        return 0.5 * (Dcell[:-1] + Dcell[1:])
    if mode == "midpoint_property":
        return moisture_diffusivity(
            0.5 * (C[:-1] + C[1:]),
            0.5 * (T_C[:-1] + T_C[1:]),
        )
    raise ValueError(f"Unknown face-average mode: {mode}")


def coupled_surface_values(
    T_C: Array,
    C: Array,
    t: float,
    grid: FVMGrid,
    env: EnvironmentData,
    params: Q2Parameters,
    config: SolverConfig,
) -> tuple[float, float, int]:
    """Solve the local coupled Robin problem for (T_s, C_s).

    Uses a Gauss-Seidel-like local Picard iteration:
        C_s^m -> k(C_s^m) -> T_s^(m+1)
        -> D(C_s^m, T_s^(m+1)) -> C_s^(m+1)
    The cell-centered fields are held fixed during this local solve.
    """
    if T_C.size < 2 or C.size < 2:
        raise ValueError("At least two cells are required.")

    Tp, Tpp = float(T_C[-1]), float(T_C[-2])
    Cp, Cpp = float(C[-1]), float(C[-2])
    Tinf = float(env.T_inf(t))
    Cinf = float(env.C_inf(t))
    dr = grid.dr

    # A good initial guess is the boundary value reconstructed with the last-cell
    # material state; then iterate the local thermo-moisture coupling.
    Cs = Cp
    ks = float(thermal_conductivity(Cs))
    Ts = (ks * (9.0 * Tp - Tpp) + 3.0 * dr * params.h * Tinf) / (
        8.0 * ks + 3.0 * dr * params.h
    )

    omega = config.surface_relaxation
    if not (0.0 < omega <= 1.0):
        raise ValueError("surface_relaxation must lie in (0, 1].")

    for it in range(1, config.surface_max_iter + 1):
        ks = float(thermal_conductivity(Cs))
        Ts_raw = (ks * (9.0 * Tp - Tpp) + 3.0 * dr * params.h * Tinf) / (
            8.0 * ks + 3.0 * dr * params.h
        )
        Ds = float(moisture_diffusivity(Cs, Ts_raw))
        Cs_raw = (Ds * (9.0 * Cp - Cpp) + 3.0 * dr * params.hm * Cinf) / (
            8.0 * Ds + 3.0 * dr * params.hm
        )
        Ts_new = omega * Ts_raw + (1.0 - omega) * Ts
        Cs_new = omega * Cs_raw + (1.0 - omega) * Cs

        if Cs_new <= 0.0 or not (math.isfinite(Ts_new) and math.isfinite(Cs_new)):
            raise FloatingPointError(
                f"Surface Picard produced an invalid state at t={t}: Ts={Ts_new}, Cs={Cs_new}"
            )

        if abs(Ts_new - Ts) <= config.surface_tol_T and abs(Cs_new - Cs) <= config.surface_tol_C:
            return float(Ts_new), float(Cs_new), it
        Ts, Cs = Ts_new, Cs_new

    raise RuntimeError(
        f"Surface Picard did not converge at t={t:.6g} s after "
        f"{config.surface_max_iter} iterations. Last Ts={Ts:.8g}, Cs={Cs:.8g}."
    )


def joint_rhs(
    t: float,
    y: Array,
    grid: FVMGrid,
    env: EnvironmentData,
    params: Q2Parameters,
    config: SolverConfig,
    surface_counter: Optional[dict] = None,
) -> Array:
    N = grid.N
    T_C = np.asarray(y[:N], dtype=float)
    C = np.asarray(y[N:], dtype=float)
    check_material_state(T_C, C)

    kf = _k_faces(C, config.face_average)
    Df = _D_faces(C, T_C, config.face_average)

    q = np.zeros(N + 1, dtype=float)  # positive in +r direction (outward)
    J = np.zeros(N + 1, dtype=float)

    q[1:N] = -kf * (T_C[1:] - T_C[:-1]) / grid.dr
    J[1:N] = -Df * (C[1:] - C[:-1]) / grid.dr

    Ts, Cs, niter = coupled_surface_values(T_C, C, t, grid, env, params, config)
    q[N] = params.h * (Ts - float(env.T_inf(t)))
    J[N] = params.hm * (Cs - float(env.C_inf(t)))

    if surface_counter is not None:
        surface_counter["calls"] = surface_counter.get("calls", 0) + 1
        surface_counter["iterations"] = surface_counter.get("iterations", 0) + niter
        surface_counter["max_iterations"] = max(surface_counter.get("max_iterations", 0), niter)

    net_q = grid.areas[:-1] * q[:-1] - grid.areas[1:] * q[1:]
    net_J = grid.areas[:-1] * J[:-1] - grid.areas[1:] * J[1:]

    cap = density(C) * heat_capacity(C)
    dTdt = net_q / (cap * grid.volumes)
    dCdt = net_J / grid.volumes

    return np.concatenate([dTdt, dCdt])


def jacobian_sparsity(N: int) -> csr_matrix:
    """Conservative block-banded sparsity for y=[T, C]."""
    rows: list[int] = []
    cols: list[int] = []
    for i in range(N):
        neigh = range(max(0, i - 1), min(N, i + 2))
        for j in neigh:
            # T row depends on neighboring T and C.
            rows.extend([i, i])
            cols.extend([j, N + j])
            # C row depends on neighboring T and C through D(C,T).
            rows.extend([N + i, N + i])
            cols.extend([j, N + j])
    data = np.ones(len(rows), dtype=bool)
    return csr_matrix((data, (rows, cols)), shape=(2 * N, 2 * N))


def center_reconstruction(u: Array) -> float:
    return float((9.0 * u[0] - u[1]) / 8.0)


def _surface_output_values(
    T_C: Array,
    C: Array,
    t: float,
    grid: FVMGrid,
    env: EnvironmentData,
    params: Q2Parameters,
    config: SolverConfig,
) -> tuple[float, float]:
    if abs(t) <= 1e-14:
        # Exact requested initial condition. The moisture Robin condition is
        # intentionally incompatible at t=0 and acts from t=0+ onward.
        return params.T0_C, params.C0
    Ts, Cs, _ = coupled_surface_values(T_C, C, t, grid, env, params, config)
    return Ts, Cs


def sample_state_pair(
    T_C: Array,
    C: Array,
    t: float,
    target_r: Array,
    grid: FVMGrid,
    env: EnvironmentData,
    params: Q2Parameters,
    config: SolverConfig,
) -> tuple[Array, Array]:
    target_r = np.asarray(target_r, dtype=float)
    if np.any(target_r < -1e-14) or np.any(target_r > grid.R + 1e-14):
        raise ValueError("Target radii must lie within [0, R].")

    T0 = center_reconstruction(T_C)
    C0 = center_reconstruction(C)
    Ts, Cs = _surface_output_values(T_C, C, t, grid, env, params, config)
    r_aug = np.concatenate([[0.0], grid.centers, [grid.R]])
    T_aug = np.concatenate([[T0], T_C, [Ts]])
    C_aug = np.concatenate([[C0], C, [Cs]])
    return np.interp(target_r, r_aug, T_aug), np.interp(target_r, r_aug, C_aug)


def augmented_snapshot(sol: Q2Solution, field: Literal["T", "C"], t: float) -> tuple[Array, Array]:
    idx = np.where(np.isclose(sol.state_t, t, atol=1e-10, rtol=0))[0]
    if idx.size == 0:
        raise ValueError("Requested t must be one of sol.state_t for an augmented full snapshot.")
    j = int(idx[0])
    Tj = sol.T_state[j]
    Cj = sol.C_state[j]
    T0 = center_reconstruction(Tj)
    C0 = center_reconstruction(Cj)
    Ts, Cs = _surface_output_values(Tj, Cj, float(t), sol.grid, sol.env, sol.params, sol.config)
    r_aug = np.concatenate([[0.0], sol.grid.centers, [sol.grid.R]])
    if field == "T":
        return r_aug, np.concatenate([[T0], Tj, [Ts]])
    if field == "C":
        return r_aug, np.concatenate([[C0], Cj, [Cs]])
    raise ValueError("field must be 'T' or 'C'.")


def _global_output_times(t_end: float, output_dt: float) -> Array:
    if output_dt <= 0:
        raise ValueError("output_dt must be positive.")
    n = int(round(t_end / output_dt))
    if not np.isclose(n * output_dt, t_end, atol=1e-10, rtol=0):
        raise ValueError("For this workflow, t_end must be an integer multiple of output_dt.")
    return np.linspace(0.0, t_end, n + 1)


def _segment_edges(env: EnvironmentData, t_end: float, segment_at_knots: bool) -> Array:
    if not segment_at_knots:
        return np.array([0.0, t_end])
    knots = env.t[(env.t > 0.0) & (env.t < t_end)]
    return np.concatenate([[0.0], knots, [t_end]])


def solve_q2(
    env_path: str | Path,
    N: int = 640,
    config: Optional[SolverConfig] = None,
    params: Optional[Q2Parameters] = None,
    t_end: float = 10800.0,
    output_dr: float = 0.001,
) -> Q2Solution:
    """Production Q2 solver: coupled MOL + BDF, segmented at data knots.

    To avoid storing 10801 x N full states, each segment is sampled immediately
    onto the required 0.1 cm output grid. Full cell-centered states are retained
    only at segment boundaries (normally every 60 s) for diagnostics.
    """
    config = config or SolverConfig()
    params = params or Q2Parameters()
    env = EnvironmentData.from_excel(env_path)
    if t_end > env.t[-1] + 1e-12:
        raise ValueError(f"Environment data end at {env.t[-1]} s, before requested t_end={t_end} s.")
    grid = FVMGrid.build(N, params.R)

    n_r = int(round(params.R / output_dr))
    if not np.isclose(n_r * output_dr, params.R, atol=1e-12, rtol=0):
        raise ValueError("R must be an integer multiple of output_dr.")
    r_out = np.linspace(0.0, params.R, n_r + 1)
    t_out = _global_output_times(t_end, config.output_dt)

    y = np.concatenate([
        np.full(N, params.T0_C, dtype=float),
        np.full(N, params.C0, dtype=float),
    ])

    edges = _segment_edges(env, t_end, config.segment_at_data_knots)
    sparsity = jacobian_sparsity(N) if config.method.upper() in {"BDF", "RADAU"} else None
    surface_counter: dict = {}

    T_out_chunks: list[Array] = []
    C_out_chunks: list[Array] = []
    t_chunks: list[Array] = []

    state_t = [0.0]
    T_state = [y[:N].copy()]
    C_state = [y[N:].copy()]

    totals = {"segments": 0, "nfev": 0, "njev": 0, "nlu": 0}

    for s in range(len(edges) - 1):
        a, b = float(edges[s]), float(edges[s + 1])
        mask = (t_out >= a - 1e-12) & (t_out <= b + 1e-12)
        tev = t_out[mask]
        if s > 0:
            tev = tev[tev > a + 1e-12]
        if tev.size == 0 or tev[-1] < b - 1e-12:
            tev = np.concatenate([tev, [b]])
        if s == 0 and (tev.size == 0 or tev[0] > a + 1e-12):
            tev = np.concatenate([[a], tev])

        rhs = lambda tt, yy: joint_rhs(tt, yy, grid, env, params, config, surface_counter)
        ivp = solve_ivp(
            rhs,
            (a, b),
            y,
            method=config.method,
            t_eval=tev,
            rtol=config.rtol,
            atol=config.atol,
            max_step=config.max_step,
            jac_sparsity=sparsity,
        )
        if not ivp.success:
            raise RuntimeError(f"BDF failed on [{a},{b}] s: {ivp.message}")

        totals["segments"] += 1
        totals["nfev"] += int(ivp.nfev)
        totals["njev"] += int(getattr(ivp, "njev", 0) or 0)
        totals["nlu"] += int(getattr(ivp, "nlu", 0) or 0)

        # Process this short segment immediately to keep memory bounded.
        segT = []
        segC = []
        for jj, tt in enumerate(ivp.t):
            Tj = ivp.y[:N, jj]
            Cj = ivp.y[N:, jj]
            Tsamp, Csamp = sample_state_pair(
                Tj, Cj, float(tt), r_out, grid, env, params, config
            )
            segT.append(Tsamp)
            segC.append(Csamp)

        t_chunks.append(ivp.t.copy())
        T_out_chunks.append(np.asarray(segT))
        C_out_chunks.append(np.asarray(segC))

        y = ivp.y[:, -1].copy()
        state_t.append(b)
        T_state.append(y[:N].copy())
        C_state.append(y[N:].copy())

    tout = np.concatenate(t_chunks)
    T_out = np.vstack(T_out_chunks)
    C_out = np.vstack(C_out_chunks)

    # The segment logic should produce the requested global output grid exactly.
    keep = np.isin(np.round(tout / config.output_dt).astype(int),
                   np.arange(len(t_out)))
    tout = tout[keep]
    T_out = T_out[keep]
    C_out = C_out[keep]
    order = np.argsort(tout)
    tout, T_out, C_out = tout[order], T_out[order], C_out[order]
    # Remove any accidental duplicate boundary times while preserving first occurrence.
    _, unique_idx = np.unique(np.round(tout, 12), return_index=True)
    unique_idx = np.sort(unique_idx)
    tout, T_out, C_out = tout[unique_idx], T_out[unique_idx], C_out[unique_idx]

    if len(tout) != len(t_out) or not np.allclose(tout, t_out, atol=1e-10, rtol=0):
        raise RuntimeError("Internal segment assembly did not reproduce the requested output time grid.")

    totals.update({
        "surface_calls": int(surface_counter.get("calls", 0)),
        "surface_total_iterations": int(surface_counter.get("iterations", 0)),
        "surface_max_iterations": int(surface_counter.get("max_iterations", 0)),
        "surface_mean_iterations": (
            float(surface_counter.get("iterations", 0)) / max(1, int(surface_counter.get("calls", 0)))
        ),
    })

    return Q2Solution(
        grid=grid, env=env, params=params, config=config,
        t=tout, r=r_out, T=T_out, C=C_out,
        state_t=np.asarray(state_t),
        T_state=np.asarray(T_state), C_state=np.asarray(C_state),
        solver_stats=totals,
    )


# -----------------------------------------------------------------------------
# Paper outputs and diagnostics
# -----------------------------------------------------------------------------

def _index_exact(values: Array, targets: Sequence[float], name: str) -> list[int]:
    idx = []
    for x in targets:
        hit = np.where(np.isclose(values, x, atol=1e-10, rtol=0))[0]
        if hit.size == 0:
            raise ValueError(f"{name}={x} is not present on the stored grid.")
        idx.append(int(hit[0]))
    return idx


def paper_tables(sol: Q2Solution) -> tuple[pd.DataFrame, pd.DataFrame]:
    times_s = np.array([1800, 3600, 5400, 7200, 9000, 10800], dtype=float)
    radii_m = np.array([0.0, 0.005, 0.010, 0.015, 0.020])
    it = _index_exact(sol.t, times_s, "time")
    ir = _index_exact(sol.r, radii_m, "radius")
    columns = ["0", "0.5", "1.0", "1.5", "2.0"]
    index = ["0.5", "1.0", "1.5", "2.0", "2.5", "3.0"]
    Tdf = pd.DataFrame(sol.T[np.ix_(it, ir)], index=index, columns=columns)
    Cdf = pd.DataFrame(sol.C[np.ix_(it, ir)], index=index, columns=columns)
    Tdf.index.name = "Time (h)"
    Cdf.index.name = "Time (h)"
    Tdf.columns.name = "Radius (cm)"
    Cdf.columns.name = "Radius (cm)"
    return Tdf, Cdf


def physical_diagnostics(sol: Q2Solution) -> dict:
    out = {
        "T_min": float(np.min(sol.T)),
        "T_max": float(np.max(sol.T)),
        "C_min": float(np.min(sol.C)),
        "C_max": float(np.max(sol.C)),
        "all_C_positive": bool(np.all(sol.C > 0)),
        "radial_T_non_decreasing_fraction": float(np.mean(np.all(np.diff(sol.T, axis=1) >= -1e-8, axis=1))),
        "radial_C_non_increasing_fraction": float(np.mean(np.all(np.diff(sol.C, axis=1) <= 1e-8, axis=1))),
    }
    Tinf = sol.env.T_inf(sol.t)
    Cinf = sol.env.C_inf(sol.t)
    out["heat_flux_direction_fraction"] = float(np.mean((Tinf - sol.T[:, -1]) >= -1e-8))
    out["moisture_flux_direction_fraction"] = float(np.mean((sol.C[:, -1] - Cinf) >= -1e-10))
    return out


def property_ranges(sol: Q2Solution) -> dict:
    T = sol.T_state
    C = sol.C_state
    rho = density(C)
    cp = heat_capacity(C)
    k = thermal_conductivity(C)
    D = moisture_diffusivity(C, T)
    alpha = k / (rho * cp)
    return {
        "rho_min": float(rho.min()), "rho_max": float(rho.max()),
        "cp_min": float(cp.min()), "cp_max": float(cp.max()),
        "k_min": float(k.min()), "k_max": float(k.max()),
        "D_min": float(D.min()), "D_max": float(D.max()),
        "alpha_min": float(alpha.min()), "alpha_max": float(alpha.max()),
    }


def boundary_residuals(sol: Q2Solution) -> dict:
    # Independent cubic derivative at R using x/dr = [0,-1/2,-3/2,-5/2].
    weights = np.array([46/15, -15/4, 5/6, -3/20], dtype=float) / sol.grid.dr
    heat_res = []
    moist_res = []
    heat_rel = []
    moist_rel = []

    for t, T, C in zip(sol.state_t, sol.T_state, sol.C_state):
        if abs(t) <= 1e-14:
            # For residual testing, apply the Robin law at 0+ rather than the
            # exact initial output convention.
            Ts, Cs, _ = coupled_surface_values(T, C, 0.0, sol.grid, sol.env, sol.params, sol.config)
        else:
            Ts, Cs, _ = coupled_surface_values(T, C, float(t), sol.grid, sol.env, sol.params, sol.config)
        dTdr = np.dot(weights, np.array([Ts, T[-1], T[-2], T[-3]]))
        dCdr = np.dot(weights, np.array([Cs, C[-1], C[-2], C[-3]]))
        lhsT = -float(thermal_conductivity(Cs)) * dTdr
        rhsT = sol.params.h * (Ts - float(sol.env.T_inf(t)))
        lhsC = -float(moisture_diffusivity(Cs, Ts)) * dCdr
        rhsC = sol.params.hm * (Cs - float(sol.env.C_inf(t)))
        rT = lhsT - rhsT
        rC = lhsC - rhsC
        heat_res.append(rT)
        moist_res.append(rC)
        heat_rel.append(abs(rT) / max(abs(lhsT), abs(rhsT), 1e-14))
        moist_rel.append(abs(rC) / max(abs(lhsC), abs(rhsC), 1e-14))

    return {
        "time": sol.state_t.copy(),
        "heat_residual": np.asarray(heat_res),
        "moisture_residual": np.asarray(moist_res),
        "heat_relative_residual": np.asarray(heat_rel),
        "moisture_relative_residual": np.asarray(moist_rel),
        "heat_max_abs": float(np.max(np.abs(heat_res))),
        "moisture_max_abs": float(np.max(np.abs(moist_res))),
        "heat_max_relative": float(np.max(heat_rel)),
        "moisture_max_relative": float(np.max(moist_rel)),
        "heat_max_abs_after_t0": float(np.max(np.abs(np.asarray(heat_res)[1:]))) if len(heat_res) > 1 else 0.0,
        "moisture_max_abs_after_t0": float(np.max(np.abs(np.asarray(moist_res)[1:]))) if len(moist_res) > 1 else 0.0,
        "heat_max_relative_after_t0": float(np.max(np.asarray(heat_rel)[1:])) if len(heat_rel) > 1 else 0.0,
        "moisture_max_relative_after_t0": float(np.max(np.asarray(moist_rel)[1:])) if len(moist_rel) > 1 else 0.0,
    }


def balance_diagnostics(sol: Q2Solution) -> dict:
    # Heat: check the semi-discrete PDE balance that is valid for the chosen
    # apparent-heat-capacity model. Do NOT reinterpret this as d/dt ∫rho cp T dV.
    heat_res = []
    heat_rel = []
    moisture_inst_res = []
    moisture_inst_rel = []

    for t, T, C in zip(sol.state_t, sol.T_state, sol.C_state):
        y = np.concatenate([T, C])
        dy = joint_rhs(float(t), y, sol.grid, sol.env, sol.params, sol.config)
        dT = dy[:sol.grid.N]
        dC = dy[sol.grid.N:]
        Ts, Cs, _ = coupled_surface_values(T, C, float(t), sol.grid, sol.env, sol.params, sol.config)
        qR = sol.params.h * (Ts - float(sol.env.T_inf(t)))
        JR = sol.params.hm * (Cs - float(sol.env.C_inf(t)))
        heat_storage = np.sum(density(C) * heat_capacity(C) * sol.grid.volumes * dT)
        moist_storage = np.sum(sol.grid.volumes * dC)
        rH = heat_storage + sol.grid.areas[-1] * qR
        rM = moist_storage + sol.grid.areas[-1] * JR
        heat_res.append(rH)
        moisture_inst_res.append(rM)
        heat_rel.append(abs(rH) / max(abs(heat_storage), abs(sol.grid.areas[-1] * qR), 1e-14))
        moisture_inst_rel.append(abs(rM) / max(abs(moist_storage), abs(sol.grid.areas[-1] * JR), 1e-14))

    # Moisture integral balance over time is valid for the effective C equation.
    M = np.sum(sol.C_state * sol.grid.volumes[None, :], axis=1)
    Csurf = sol.C[:, -1].copy()
    # At t=0 use the Robin 0+ surface value for flux integration, not the exact
    # initial output value C0 retained for the requested result table.
    _, Cs0, _ = coupled_surface_values(
        sol.C_state[0] * 0 + sol.params.T0_C,
        sol.C_state[0], 0.0, sol.grid, sol.env, sol.params, sol.config
    )
    Csurf[0] = Cs0
    Jout = sol.params.hm * (Csurf - sol.env.C_inf(sol.t))
    cumulative_out = cumulative_trapezoid(sol.grid.areas[-1] * Jout, sol.t, initial=0.0)
    state_idx = _index_exact(sol.t, sol.state_t, "state time")
    predicted_M = M[0] - cumulative_out[state_idx]
    mass_residual = M - predicted_M
    denom = np.maximum(np.maximum(np.abs(M - M[0]), np.abs(cumulative_out[state_idx])), 1e-14)
    mass_relative = np.abs(mass_residual) / denom
    if len(mass_relative) > 1:
        mass_rel_report = float(np.max(mass_relative[1:]))
    else:
        mass_rel_report = 0.0
    final_scale = max(abs(M[-1] - M[0]), abs(cumulative_out[state_idx][-1]), 1e-14)
    final_mass_relative = float(abs(mass_residual[-1]) / final_scale)
    global_scale = max(float(np.max(np.abs(M - M[0]))), float(np.max(np.abs(cumulative_out[state_idx]))), 1e-14)
    global_mass_relative = float(np.max(np.abs(mass_residual)) / global_scale)

    return {
        "time": sol.state_t.copy(),
        "heat_balance_residual": np.asarray(heat_res),
        "heat_balance_relative": np.asarray(heat_rel),
        "moisture_instantaneous_residual": np.asarray(moisture_inst_res),
        "moisture_instantaneous_relative": np.asarray(moisture_inst_rel),
        "moisture_integral": M,
        "moisture_cumulative_out": cumulative_out[state_idx],
        "moisture_integral_residual": mass_residual,
        "moisture_integral_relative": mass_relative,
        "heat_max_relative": float(np.max(heat_rel)),
        "moisture_instantaneous_max_relative": float(np.max(moisture_inst_rel)),
        "moisture_integral_max_relative_excluding_t0": mass_rel_report,
        "moisture_integral_final_relative": final_mass_relative,
        "moisture_integral_global_normalized_max": global_mass_relative,
    }


def compare_solutions(a: Q2Solution, b: Q2Solution) -> dict:
    if not np.allclose(a.t, b.t) or not np.allclose(a.r, b.r):
        raise ValueError("Solutions must share the same output grids.")
    dT = np.abs(a.T - b.T)
    dC = np.abs(a.C - b.C)
    iT = np.unravel_index(np.argmax(dT), dT.shape)
    iC = np.unravel_index(np.argmax(dC), dC.shape)
    return {
        "max_abs_T": float(dT[iT]),
        "max_abs_C": float(dC[iC]),
        "T_location": (float(a.t[iT[0]]), float(a.r[iT[1]])),
        "C_location": (float(a.t[iC[0]]), float(a.r[iC[1]])),
    }


def grid_convergence_study(
    env_path: str | Path,
    N_values: Sequence[int] = (320, 640, 1280),
    config: Optional[SolverConfig] = None,
    params: Optional[Q2Parameters] = None,
    t_end: float = 10800.0,
) -> tuple[pd.DataFrame, dict[int, Q2Solution]]:
    config = config or SolverConfig()
    params = params or Q2Parameters()
    sols: dict[int, Q2Solution] = {}
    for N in N_values:
        sols[int(N)] = solve_q2(env_path, N=int(N), config=config, params=params, t_end=t_end)
    rows = []
    for Na, Nb in zip(N_values[:-1], N_values[1:]):
        cmp = compare_solutions(sols[int(Na)], sols[int(Nb)])
        rows.append({
            "coarse_N": int(Na), "fine_N": int(Nb),
            "max_abs_T_degC": cmp["max_abs_T"],
            "max_abs_C_kgkg": cmp["max_abs_C"],
            "T_location_t_s": cmp["T_location"][0],
            "T_location_r_cm": 100.0 * cmp["T_location"][1],
            "C_location_t_s": cmp["C_location"][0],
            "C_location_r_cm": 100.0 * cmp["C_location"][1],
        })
    return pd.DataFrame(rows), sols


def face_average_sensitivity(
    env_path: str | Path,
    N: int = 320,
    modes: Sequence[FaceAverage] = ("harmonic", "midpoint_property", "arithmetic"),
    base_config: Optional[SolverConfig] = None,
    t_end: float = 10800.0,
) -> tuple[pd.DataFrame, dict[str, Q2Solution]]:
    base = base_config or SolverConfig()
    sols: dict[str, Q2Solution] = {}
    for mode in modes:
        cfg = SolverConfig(**{**base.__dict__, "face_average": mode})
        sols[mode] = solve_q2(env_path, N=N, config=cfg, t_end=t_end)
    ref = sols[modes[0]]
    rows = []
    for mode in modes[1:]:
        cmp = compare_solutions(ref, sols[mode])
        rows.append({"reference": modes[0], "comparison": mode, **cmp})
    return pd.DataFrame(rows), sols


# -----------------------------------------------------------------------------
# Manual backward-Euler + block Picard reference solver
# -----------------------------------------------------------------------------

def _be_diffusion_solve(
    u_old: Array,
    capacity: Array,
    face_coef: Array,
    boundary_coef: float,
    transfer_coef: float,
    ambient: float,
    grid: FVMGrid,
    dt: float,
) -> Array:
    """One linear backward-Euler diffusion solve with frozen coefficients.

    The second-order Robin reconstruction is analytically folded into the last
    finite-volume row. `capacity` is rho*cp for heat and 1 for moisture.
    """
    N = grid.N
    S = np.asarray(capacity, dtype=float) * grid.volumes / dt
    g = grid.areas[1:N] * face_coef / grid.dr

    diagv = S.copy()
    lower = -g.copy()
    upper = -g.copy()
    diagv[:-1] += g
    diagv[1:] += g

    gamma = transfer_coef * boundary_coef / (8.0 * boundary_coef + 3.0 * grid.dr * transfer_coef)
    AR = grid.areas[-1]
    diagv[-1] += 9.0 * AR * gamma
    lower[-1] -= AR * gamma

    rhs = S * u_old
    rhs[-1] += 8.0 * AR * gamma * ambient

    A = diags([lower, diagv, upper], offsets=[-1, 0, 1], format="csr")
    return np.asarray(spsolve(A, rhs), dtype=float)


def solve_q2_block_picard(
    env_path: str | Path,
    N: int = 160,
    config: Optional[PicardConfig] = None,
    params: Optional[Q2Parameters] = None,
    t_end: float = 10800.0,
) -> BlockPicardSolution:
    """Independent fixed-step backward-Euler + block-Picard reference.

    This is not the production solver. Its purpose is error/splitting analysis.
    Set one_pass=True to intentionally perform one lagged T->C sweep per step.
    """
    config = config or PicardConfig()
    params = params or Q2Parameters()
    env = EnvironmentData.from_excel(env_path)
    if t_end > env.t[-1] + 1e-12:
        raise ValueError("Environment data do not cover the requested interval.")
    nsteps = int(round(t_end / config.dt))
    if not np.isclose(nsteps * config.dt, t_end, atol=1e-10, rtol=0):
        raise ValueError("t_end must be an integer multiple of Picard dt.")
    if not (0 < config.relaxation <= 1):
        raise ValueError("Picard relaxation must lie in (0,1].")

    grid = FVMGrid.build(N, params.R)
    times = np.linspace(0.0, t_end, nsteps + 1)
    T_hist = np.empty((nsteps + 1, N), dtype=float)
    C_hist = np.empty((nsteps + 1, N), dtype=float)
    iterations = np.zeros(nsteps, dtype=int)
    T_hist[0] = params.T0_C
    C_hist[0] = params.C0

    surf_cfg = SolverConfig(face_average=config.face_average)

    for n in range(nsteps):
        tnew = float(times[n + 1])
        Told = T_hist[n]
        Cold = C_hist[n]
        Tg = Told.copy()
        Cg = Cold.copy()

        converged = False
        maxit = 1 if config.one_pass else config.max_iter
        for m in range(1, maxit + 1):
            check_material_state(Tg, Cg)
            # Frozen thermal coefficients from the current moisture iterate.
            Ts_g, Cs_g, _ = coupled_surface_values(Tg, Cg, tnew, grid, env, params, surf_cfg)
            cap = density(Cg) * heat_capacity(Cg)
            kf = _k_faces(Cg, config.face_average)
            ks = float(thermal_conductivity(Cs_g))
            Tlin = _be_diffusion_solve(
                Told, cap, kf, ks, params.h, float(env.T_inf(tnew)), grid, config.dt
            )
            Tnew = config.relaxation * Tlin + (1.0 - config.relaxation) * Tg

            # Frozen moisture diffusivity from current C iterate and new T iterate.
            Ts_aux, Cs_aux, _ = coupled_surface_values(Tnew, Cg, tnew, grid, env, params, surf_cfg)
            Df = _D_faces(Cg, Tnew, config.face_average)
            Ds = float(moisture_diffusivity(Cs_aux, Ts_aux))
            Clin = _be_diffusion_solve(
                Cold, np.ones(N), Df, Ds, params.hm, float(env.C_inf(tnew)), grid, config.dt
            )
            Cnew = config.relaxation * Clin + (1.0 - config.relaxation) * Cg

            if np.any(Cnew <= 0) or not np.all(np.isfinite(Cnew)) or not np.all(np.isfinite(Tnew)):
                raise FloatingPointError(f"Invalid block-Picard state at t={tnew} s.")

            eT = float(np.max(np.abs(Tnew - Tg)))
            eC = float(np.max(np.abs(Cnew - Cg)))
            Tg, Cg = Tnew, Cnew

            if config.one_pass or (eT <= config.tol_T and eC <= config.tol_C):
                converged = True
                iterations[n] = m
                break

        if not converged:
            raise RuntimeError(
                f"Block Picard failed at t={tnew:.3f} s after {config.max_iter} iterations. "
                f"Try a smaller dt or under-relaxation."
            )

        T_hist[n + 1] = Tg
        C_hist[n + 1] = Cg

    return BlockPicardSolution(
        grid=grid, env=env, params=params, config=config,
        t=times, T_state=T_hist, C_state=C_hist, iterations=iterations,
    )


def sample_block_picard(
    sol: BlockPicardSolution,
    times: Sequence[float],
    radii: Sequence[float],
) -> tuple[Array, Array]:
    times = np.asarray(times, dtype=float)
    radii = np.asarray(radii, dtype=float)
    cfg = SolverConfig(face_average=sol.config.face_average)
    Tout = np.empty((len(times), len(radii)))
    Cout = np.empty_like(Tout)

    for it, tt in enumerate(times):
        if tt < sol.t[0] - 1e-12 or tt > sol.t[-1] + 1e-12:
            raise ValueError("Requested time outside Picard solution interval.")
        hi = int(np.searchsorted(sol.t, tt, side="left"))
        if hi < len(sol.t) and np.isclose(sol.t[hi], tt, atol=1e-10, rtol=0):
            Tcell = sol.T_state[hi]
            Ccell = sol.C_state[hi]
        else:
            lo = hi - 1
            w = (tt - sol.t[lo]) / (sol.t[hi] - sol.t[lo])
            Tcell = (1 - w) * sol.T_state[lo] + w * sol.T_state[hi]
            Ccell = (1 - w) * sol.C_state[lo] + w * sol.C_state[hi]
        Tout[it], Cout[it] = sample_state_pair(
            Tcell, Ccell, float(tt), radii, sol.grid, sol.env, sol.params, cfg
        )
    return Tout, Cout


def compare_bdf_picard(
    bdf_sol: Q2Solution,
    picard_sol: BlockPicardSolution,
    times: Sequence[float] = (1800, 3600, 5400, 7200, 9000, 10800),
    radii: Sequence[float] = (0.0, 0.005, 0.010, 0.015, 0.020),
) -> dict:
    times = np.asarray(times, dtype=float)
    radii = np.asarray(radii, dtype=float)
    it = _index_exact(bdf_sol.t, times, "time")
    ir = _index_exact(bdf_sol.r, radii, "radius")
    Tb = bdf_sol.T[np.ix_(it, ir)]
    Cb = bdf_sol.C[np.ix_(it, ir)]
    Tp, Cp = sample_block_picard(picard_sol, times, radii)
    dT = np.abs(Tb - Tp)
    dC = np.abs(Cb - Cp)
    return {
        "max_abs_T": float(np.max(dT)),
        "max_abs_C": float(np.max(dC)),
        "mean_abs_T": float(np.mean(dT)),
        "mean_abs_C": float(np.mean(dC)),
        "T_diff": dT,
        "C_diff": dC,
    }


def picard_error_study(
    env_path: str | Path,
    N: int = 160,
    dt_values: Sequence[float] = (60.0, 30.0, 15.0),
    bdf_config: Optional[SolverConfig] = None,
    t_end: float = 10800.0,
) -> tuple[pd.DataFrame, Q2Solution, dict[float, BlockPicardSolution]]:
    bcfg = bdf_config or SolverConfig(max_step=5.0)
    bdf = solve_q2(env_path, N=N, config=bcfg, t_end=t_end)
    rows = []
    pics: dict[float, BlockPicardSolution] = {}
    for dt in dt_values:
        pcfg = PicardConfig(dt=float(dt), one_pass=False)
        psol = solve_q2_block_picard(env_path, N=N, config=pcfg, t_end=t_end)
        pics[float(dt)] = psol
        cmp = compare_bdf_picard(bdf, psol)
        rows.append({
            "dt_s": float(dt),
            "max_abs_T_vs_BDF": cmp["max_abs_T"],
            "max_abs_C_vs_BDF": cmp["max_abs_C"],
            "mean_picard_iterations": float(np.mean(psol.iterations)),
            "max_picard_iterations": int(np.max(psol.iterations)),
        })
    return pd.DataFrame(rows), bdf, pics


def splitting_error_study(
    env_path: str | Path,
    N: int = 160,
    dt: float = 30.0,
    t_end: float = 10800.0,
) -> tuple[pd.DataFrame, BlockPicardSolution, BlockPicardSolution]:
    converged = solve_q2_block_picard(
        env_path, N=N, config=PicardConfig(dt=dt, one_pass=False), t_end=t_end
    )
    onepass = solve_q2_block_picard(
        env_path, N=N, config=PicardConfig(dt=dt, one_pass=True), t_end=t_end
    )
    times = np.array([1800, 3600, 5400, 7200, 9000, 10800], dtype=float)
    radii = np.array([0.0, 0.005, 0.010, 0.015, 0.020])
    Tc, Cc = sample_block_picard(converged, times, radii)
    To, Co = sample_block_picard(onepass, times, radii)
    df = pd.DataFrame([{
        "N": N,
        "dt_s": dt,
        "max_abs_T_onepass_vs_converged": float(np.max(np.abs(To - Tc))),
        "max_abs_C_onepass_vs_converged": float(np.max(np.abs(Co - Cc))),
        "mean_converged_picard_iterations": float(np.mean(converged.iterations)),
        "max_converged_picard_iterations": int(np.max(converged.iterations)),
    }])
    return df, converged, onepass


# -----------------------------------------------------------------------------
# Result writer
# -----------------------------------------------------------------------------

def write_result2(
    sol: Q2Solution,
    output_path: str | Path,
    template_path: Optional[str | Path] = None,
    decimals: int = 4,
) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    if template_path is not None:
        wb = load_workbook(template_path)
    else:
        wb = Workbook()
        default = wb.active
        wb.remove(default)
        wb.create_sheet("温度")
        wb.create_sheet("水分浓度")

    for sheet_name, values in [("温度", sol.T), ("水分浓度", sol.C)]:
        if sheet_name not in wb.sheetnames:
            wb.create_sheet(sheet_name)
        ws = wb[sheet_name]
        ws.cell(row=1, column=1, value="时间/s")
        for j, r in enumerate(sol.r, start=2):
            ws.cell(row=1, column=j, value=round(float(r * 100.0), 10))
        for i, tt in enumerate(sol.t, start=2):
            ws.cell(row=i, column=1, value=int(round(float(tt))))
            for j in range(values.shape[1]):
                ws.cell(row=i, column=j + 2, value=round(float(values[i - 2, j]), decimals))

    wb.save(output_path)
    return output_path


def explicit_stability_reference(N: int, params: Optional[Q2Parameters] = None) -> dict:
    params = params or Q2Parameters()
    grid = FVMGrid.build(N, params.R)
    C0 = params.C0
    alpha0 = float(thermal_conductivity(C0) / (density(C0) * heat_capacity(C0)))
    D0 = float(moisture_diffusivity(C0, params.T0_C))
    # Conservative radial explicit-Euler references, not production limits.
    return {
        "N": N,
        "dr_m": grid.dr,
        "dr_cm": 100.0 * grid.dr,
        "alpha0_m2s": alpha0,
        "D0_m2s": D0,
        "heat_dt_reference_s": 0.25 * grid.dr**2 / alpha0,
        "moisture_dt_reference_s": 0.25 * grid.dr**2 / D0,
    }


# -----------------------------------------------------------------------------
# CLI
# -----------------------------------------------------------------------------

def _build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(description="Q2 drying solver: cylindrical FVM + coupled BDF")
    p.add_argument("--env", required=True, help="Path to Attachment 1 xlsx")
    p.add_argument("--output", required=True, help="Path to result2.xlsx")
    p.add_argument("--template", default=None, help="Optional official result2.xlsx template")
    p.add_argument("--N", type=int, default=640, help="Number of FVM cells")
    p.add_argument("--method", default="BDF")
    p.add_argument("--rtol", type=float, default=1e-9)
    p.add_argument("--atol", type=float, default=1e-11)
    p.add_argument("--max-step", type=float, default=5.0)
    p.add_argument("--output-dt", type=float, default=1.0)
    p.add_argument("--face-average", choices=["harmonic", "arithmetic", "midpoint_property"], default="harmonic")
    p.add_argument("--no-segment", action="store_true", help="Do not restart BDF at environment-data knots")
    return p


def main() -> None:
    args = _build_parser().parse_args()
    cfg = SolverConfig(
        method=args.method, rtol=args.rtol, atol=args.atol,
        max_step=args.max_step, output_dt=args.output_dt,
        face_average=args.face_average,
        segment_at_data_knots=not args.no_segment,
    )
    sol = solve_q2(args.env, N=args.N, config=cfg)
    print("Solver stats:", sol.solver_stats)
    print("Physical diagnostics:", physical_diagnostics(sol))
    print("Property ranges:", property_ranges(sol))
    print("Paper temperature table:\n", paper_tables(sol)[0])
    print("Paper moisture table:\n", paper_tables(sol)[1])
    out = write_result2(sol, args.output, template_path=args.template)
    print("Saved:", out)


if __name__ == "__main__":
    main()
